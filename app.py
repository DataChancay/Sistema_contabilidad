#! .\venv_ofi\Scripts\python.exe
from __future__ import annotations

import base64
import io
import json
import os
import re
import time as time_module
import unicodedata
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from urllib.parse import quote, urlencode
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import requests
from auth import has_permission, init_auth, permission_required
from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request, send_file, send_from_directory, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry


BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

app = Flask(__name__)
app.config.update(
    JSON_AS_ASCII=False,
    MAX_CONTENT_LENGTH=64 * 1024 * 1024,
    SEND_FILE_MAX_AGE_DEFAULT=0,
    TEMPLATES_AUTO_RELOAD=True,
)
app.jinja_env.auto_reload = True
init_auth(app)

XAFIRO_BASE_URL = os.getenv("XAFIRO_BASE_URL", "https://hotel.xafiro.net").rstrip("/")
NEO_AUTH_BASE_URL = os.getenv("NEO_AUTH_BASE_URL", "https://neo.castillodechancay.com/ms_auth").rstrip("/")
NEO_BOLETERIA_BASE_URL = os.getenv(
    "NEO_BOLETERIA_BASE_URL",
    "https://neo.castillodechancay.com/ms_boleteria",
).rstrip("/")
CULQI_API_BASE_URL = os.getenv(
    "CULQI_API_BASE_URL",
    "https://api.panel.culqi.com/or-channel-panel/api",
).rstrip("/")
MIFACT_LOGIN_BASE_URL = os.getenv("MIFACT_LOGIN_BASE_URL", "https://api.mifact.net/loginws").rstrip("/")
MIFACT_LOGIN_BASIC_USER = os.getenv("MIFACT_LOGIN_BASIC_USER", "MifactLogin2020")
MIFACT_LOGIN_BASIC_PASSWORD = os.getenv("MIFACT_LOGIN_BASIC_PASSWORD", "LoginMifact")


def resolve_culqi_timezone() -> timezone:
    timezone_name = os.getenv("CULQI_TIMEZONE", "America/Lima").strip() or "America/Lima"
    try:
        return ZoneInfo(timezone_name)
    except ZoneInfoNotFoundError:
        if timezone_name in {"America/Lima", "-05:00", "UTC-05:00"}:
            return timezone(timedelta(hours=-5), "America/Lima")
        raise


CULQI_TIMEZONE = resolve_culqi_timezone()
REQUEST_TIMEOUT = (20, 180)
CULQI_CURRENCIES = ["PEN", "USD"]
CULQI_CARD_TYPES = ["CARD", "EF"]
CULQI_CARD_BRANDS = ["04", "05", "03", "07", "08", "09", "99", "97"]
CULQI_PARTIAL_PAYMENTS = ["CR", "DB", "PRE"]
CULQI_SALES_STATUS = ["01", "02", "03", "04", "05", "06", "07", "08", "09", "10", "11", "12"]
CULQI_WALLETS = ["01", "02", "03"]
CULQI_EMPTY_EXPORT_HEADERS = [
    "Empresa",
    "Comercio",
    "Producto",
    "ID Venta",
    "Marca",
    "Nro. Tarjeta",
    "Ult. 4 digitos",
    "Moneda",
    "Fecha de la transaccion",
    "Hora de la transaccion",
    "Nombres",
    "Apellidos",
    "Correo Electronico",
    "Pais",
    "Ciudad",
    "Direccion",
    "Telefono",
    "Nombre Banco",
    "Pais Banco",
    "Codigo Referencia",
    "Codigo Autorizacion",
    "ID Terminal",
    "Devolucion",
    "Pre-autorizacion",
    "Monto VENTA",
    "Venta Final",
    "Comision Emisor",
    "Comision Culqi",
    "IGV Emisor",
    "IGV Culqi",
    "IGV TOTAL",
    "Comision TOTAL",
    "Monto Aproximado Abono",
    "ID Transaccion",
    "Serie Terminal",
    "Monto Propina",
    "Estado",
    "Categoria de Estado",
    "Mensaje al Comercio",
    "Mensaje al Usuario",
    "Modo de pago",
    "Marca QR",
    "Tipo de Pago",
    "Metadata",
    "Nro Pedido",
    "Nro orden PE",
    "Fecha expiracion PE",
    "ID Comercio",
    "Lote",
    "Ref. Lote",
    "Descripcion",
    "Tipo Tokenizacion ",
    "Codigo Interno",
    "Comision + IGV Culqi",
]
MIFACT_DOCUMENT_TYPES = [
    {"id": "01", "descripcion": "Factura", "nomDoc": "", "flgEstado": True, "codModulo": 0},
    {"id": "03", "descripcion": "Boleta de venta", "nomDoc": "", "flgEstado": True, "codModulo": 0},
    {"id": "07", "descripcion": "Nota de crédito", "nomDoc": "", "flgEstado": True, "codModulo": 0},
    {"id": "08", "descripcion": "Nota de débito", "nomDoc": "", "flgEstado": True, "codModulo": 0},
]
MIFACT_SOURCES = [
    {
        "key": "asociacion",
        "name": "Asociación Civil Castillo de Chancay",
        "url": os.getenv(
            "MIFACT_ASOCIACION_URL",
            "https://api5.mifact.net/rapirest/api/vista/exportarExcelDocumento",
        ),
        "token_envs": (
            "MIFACT_ASOCIACION_TOKEN",
            "MIFACT_ASOCIACION_AUTHORIZATION",
            "MIFACT_ASOCIACION_BEARER",
            "MIFACT_TOKEN_ASOCIACION",
            "MIFACT_TOKEN_ASOCIACION_CIVIL",
            "Mifact_asociacion_token",
            "Mifact_token_asociacion",
            "Mifact_token_asociacion_civil",
        ),
        "numero_documento_emisor": "20601172420",
    },
    {
        "key": "resort",
        "name": "Resort Chancay JL S.A.C.",
        "url": os.getenv(
            "MIFACT_RESORT_URL",
            "https://api6.mifact.net/uapirest/api/vista/exportarExcelDocumento",
        ),
        "token_envs": (
            "MIFACT_RESORT_TOKEN",
            "MIFACT_RESORT_AUTHORIZATION",
            "MIFACT_RESORT_BEARER",
            "MIFACT_TOKEN_RESORT",
            "MIFACT_TOKEN_RESORT_JL",
            "Mifact_resort_token",
            "Mifact_token_resort",
            "Mifact_token_resort_jl",
        ),
        "numero_documento_emisor": "20555523042",
    },
]


class XafiroError(RuntimeError):
    """Error controlado al comunicarse con Xafiro."""


class NeoError(RuntimeError):
    """Error controlado al comunicarse con Neo."""


class CulqiError(RuntimeError):
    """Error controlado al comunicarse con Culqi."""


class MifactError(RuntimeError):
    """Error controlado al comunicarse con Mifact."""


class ConsolidadoError(RuntimeError):
    """Error controlado al procesar archivos para consolidado."""


def static_asset(filename: str) -> str:
    asset_path = BASE_DIR / "static" / filename
    version = int(asset_path.stat().st_mtime) if asset_path.exists() else 0
    return url_for("static", filename=filename, v=version)


def animation_asset(filename: str) -> str:
    asset_path = BASE_DIR / "animaciones" / filename
    version = int(asset_path.stat().st_mtime) if asset_path.exists() else 0
    return url_for("animation_file", filename=filename, v=version)


app.jinja_env.globals["static_asset"] = static_asset
app.jinja_env.globals["animation_asset"] = animation_asset
app.jinja_env.globals["has_permission"] = has_permission


def xafiro_request_error(exc: requests.RequestException, action: str) -> XafiroError:
    detail = str(exc)

    if "WinError 10013" in detail:
        return XafiroError(
            f"La PC bloqueó la conexión de Python hacia Xafiro al {action}. "
            "Revisa firewall, antivirus o permisos de red para Python."
        )
    if isinstance(exc, requests.Timeout):
        return XafiroError(
            f"Xafiro tardó demasiado en responder al {action}. "
            "Intenta nuevamente o usa un rango de fechas más corto."
        )
    if isinstance(exc, requests.ConnectionError):
        return XafiroError(
            f"No se pudo conectar con Xafiro al {action}. "
            "Verifica internet, VPN/proxy o que https://hotel.xafiro.net esté disponible."
        )
    if isinstance(exc, requests.HTTPError) and exc.response is not None:
        return XafiroError(
            f"Xafiro respondió con error HTTP {exc.response.status_code} al {action}."
        )

    return XafiroError(f"No se pudo completar la comunicación con Xafiro al {action}.")


def build_xafiro_session() -> requests.Session:
    retry_strategy = Retry(
        total=2,
        connect=2,
        read=2,
        backoff_factor=0.7,
        status_forcelist=(502, 503, 504),
        allowed_methods=frozenset({"GET", "POST"}),
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry_strategy)

    session = requests.Session()
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    session.headers.update(
        {
            "User-Agent": "Contabilidad-Chancay/0.1",
            "Accept": "application/vnd.ms-excel, application/vnd.openxmlformats-officedocument.spreadsheetml.sheet, application/octet-stream",
        }
    )
    return session


def neo_request_error(exc: requests.RequestException, action: str) -> NeoError:
    detail = str(exc)

    if "WinError 10013" in detail:
        return NeoError(
            f"La PC bloqueó la conexión de Python hacia Neo al {action}. "
            "Revisa firewall, antivirus o permisos de red para Python."
        )
    if isinstance(exc, requests.Timeout):
        return NeoError(
            f"Neo tardó demasiado en responder al {action}. "
            "Intenta nuevamente o usa un rango de fechas más corto."
        )
    if isinstance(exc, requests.ConnectionError):
        return NeoError(
            f"No se pudo conectar con Neo al {action}. "
            "Verifica internet, VPN/proxy o que https://neo.castillodechancay.com esté disponible."
        )
    if isinstance(exc, requests.HTTPError) and exc.response is not None:
        return NeoError(f"Neo respondió con error HTTP {exc.response.status_code} al {action}.")

    return NeoError(f"No se pudo completar la comunicación con Neo al {action}.")


def build_neo_session() -> requests.Session:
    retry_strategy = Retry(
        total=2,
        connect=2,
        read=2,
        backoff_factor=0.7,
        status_forcelist=(502, 503, 504),
        allowed_methods=frozenset({"GET", "POST"}),
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry_strategy)

    session = requests.Session()
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    session.headers.update(
        {
            "User-Agent": "Contabilidad-Chancay/0.1",
            "Accept": "application/json, application/vnd.openxmlformats-officedocument.spreadsheetml.sheet, application/octet-stream",
        }
    )
    return session


def culqi_request_error(exc: requests.RequestException, action: str) -> CulqiError:
    detail = str(exc)

    if "WinError 10013" in detail:
        return CulqiError(
            f"La PC bloqueó la conexión de Python hacia Culqi al {action}. "
            "Revisa firewall, antivirus o permisos de red para Python."
        )
    if isinstance(exc, requests.Timeout):
        return CulqiError(
            f"Culqi tardó demasiado en responder al {action}. "
            "Intenta nuevamente o usa un rango de fechas más corto."
        )
    if isinstance(exc, requests.ConnectionError):
        return CulqiError(
            f"No se pudo conectar con Culqi al {action}. "
            "Verifica internet, VPN/proxy o que https://culqipanel.culqi.com esté disponible."
        )
    if isinstance(exc, requests.HTTPError) and exc.response is not None:
        return CulqiError(f"Culqi respondió con error HTTP {exc.response.status_code} al {action}.")

    return CulqiError(f"No se pudo completar la comunicación con Culqi al {action}.")


def build_culqi_session() -> requests.Session:
    retry_strategy = Retry(
        total=2,
        connect=2,
        read=2,
        backoff_factor=0.7,
        status_forcelist=(502, 503, 504),
        allowed_methods=frozenset({"GET", "POST"}),
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry_strategy)

    session = requests.Session()
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    session.headers.update(
        {
            "User-Agent": "Mozilla/5.0 Contabilidad-Chancay/0.1",
            "Accept": "application/json, text/plain, */*",
            "Content-Type": "application/json",
            "Origin": "https://culqipanel.culqi.com",
            "Referer": "https://culqipanel.culqi.com/",
            "X-CULQI-ENV": "live",
        }
    )
    return session


def mifact_request_error(exc: requests.RequestException, action: str) -> MifactError:
    detail = str(exc)

    if "WinError 10013" in detail:
        return MifactError(
            f"La PC bloqueó la conexión de Python hacia Mifact al {action}. "
            "Revisa firewall, antivirus o permisos de red para Python."
        )
    if isinstance(exc, requests.Timeout):
        return MifactError(
            f"Mifact tardó demasiado en responder al {action}. "
            "Intenta nuevamente o usa un rango de fechas más corto."
        )
    if isinstance(exc, requests.ConnectionError):
        return MifactError(
            f"No se pudo conectar con Mifact al {action}. "
            "Verifica internet, VPN/proxy o que Mifact esté disponible."
        )
    if isinstance(exc, requests.HTTPError) and exc.response is not None:
        if exc.response.status_code in {401, 403}:
            return MifactError(
                f"Mifact rechazó la autorización al {action}. "
                "La cuenta configurada no tiene acceso a ese RUC o falta el token Bearer de esa empresa en .env."
            )
        return MifactError(f"Mifact respondió con error HTTP {exc.response.status_code} al {action}.")

    return MifactError(f"No se pudo completar la comunicación con Mifact al {action}.")


def build_mifact_session() -> requests.Session:
    retry_strategy = Retry(
        total=2,
        connect=2,
        read=2,
        backoff_factor=0.7,
        status_forcelist=(502, 503, 504),
        allowed_methods=frozenset({"POST"}),
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry_strategy)

    session = requests.Session()
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    session.headers.update(
        {
            "User-Agent": "Mozilla/5.0 Contabilidad-Chancay/0.1",
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet, application/vnd.ms-excel, application/octet-stream, application/json",
            "Content-Type": "application/json",
        }
    )
    return session


def parse_iso_date(raw_value: object, field_name: str) -> date:
    if not isinstance(raw_value, str):
        raise ValueError(f"{field_name} es obligatorio.")

    try:
        return datetime.strptime(raw_value, "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError(f"{field_name} debe tener el formato AAAA-MM-DD.") from exc


def xafiro_credentials() -> dict[str, str]:
    mapping = {
        "empresa": os.getenv("XAFIRO_EMPRESA", "").strip(),
        "usuario": os.getenv("XAFIRO_USUARIO", "").strip(),
        "pass": os.getenv("XAFIRO_PASSWORD", "").strip(),
    }
    if not all(mapping.values()):
        raise XafiroError("Faltan credenciales de Xafiro en el archivo .env.")
    return mapping


def neo_credentials() -> dict[str, str]:
    mapping = {
        "usuario": (
            os.getenv("NEO_USUARIO", "").strip()
            or os.getenv("Neo_usuario", "").strip()
        ),
        "password": (
            os.getenv("NEO_PASSWORD", "").strip()
            or os.getenv("Neo_password", "").strip()
        ),
    }
    if not all(mapping.values()):
        raise NeoError("Faltan credenciales de Neo en el archivo .env.")
    return mapping


def culqi_credentials() -> dict[str, str]:
    mapping = {
        "email": (
            os.getenv("CULQI_USUARIO", "").strip()
            or os.getenv("CULQUI_USUARIO", "").strip()
            or os.getenv("Culqui_usuario", "").strip()
        ),
        "password": (
            os.getenv("CULQI_PASSWORD", "").strip()
            or os.getenv("CULQUI_PASSWORD", "").strip()
            or os.getenv("Culqui_password", "").strip()
        ),
    }
    if not all(mapping.values()):
        raise CulqiError("Faltan credenciales de Culqi en el archivo .env.")
    return mapping


def mifact_authorization(source: dict[str, object]) -> str:
    token_envs = source.get("token_envs")
    if not isinstance(token_envs, tuple):
        raise MifactError("La configuración interna de Mifact es inválida.")

    token = ""
    for env_name in token_envs:
        token = os.getenv(str(env_name), "").strip()
        if token:
            break

    if not token:
        raise MifactError(
            f"Falta el token de Mifact para {source.get('name') or 'una empresa'} en el archivo .env."
        )

    return token if token.lower().startswith("bearer ") else f"Bearer {token}"


def mifact_credentials() -> dict[str, str]:
    mapping = {
        "usuario": (
            os.getenv("MIFACT_USUARIO", "").strip()
            or os.getenv("Mifact_usuario", "").strip()
        ),
        "password": (
            os.getenv("MIFACT_PASSWORD", "").strip()
            or os.getenv("Mifact_password", "").strip()
        ),
    }
    if not all(mapping.values()):
        raise MifactError("Faltan las credenciales de usuario/clave de Mifact en el archivo .env.")
    return mapping


def login_to_mifact(session: requests.Session) -> str:
    credentials = mifact_credentials()
    basic = base64.b64encode(
        f"{MIFACT_LOGIN_BASIC_USER}:{MIFACT_LOGIN_BASIC_PASSWORD}".encode("utf-8")
    ).decode("ascii")

    try:
        response = session.post(
            f"{MIFACT_LOGIN_BASE_URL}/oauth/token",
            data=urlencode(
                {
                    "grant_type": "password",
                    "username": credentials["usuario"],
                    "password": credentials["password"],
                }
            ),
            headers={
                "Content-Type": "application/x-www-form-urlencoded",
                "Authorization": f"Basic {basic}",
                "Accept": "application/json",
            },
            timeout=REQUEST_TIMEOUT,
        )
        response.raise_for_status()
    except requests.RequestException as exc:
        raise mifact_request_error(exc, "iniciar sesión") from exc

    try:
        payload = response.json()
    except ValueError as exc:
        raise MifactError("Mifact devolvió una respuesta de login inválida.") from exc

    access_token = str(payload.get("access_token") or "").strip()
    token_type = str(payload.get("token_type") or "Bearer").strip() or "Bearer"
    if not access_token:
        message = payload.get("error_description") or payload.get("message") or payload.get("error")
        raise MifactError(str(message or "Mifact no devolvió un token de acceso válido."))

    token = f"{token_type} {access_token}".strip()
    session.headers["Authorization"] = token
    return token


def resolve_mifact_authorization(session: requests.Session, source: dict[str, object]) -> str:
    try:
        return mifact_authorization(source)
    except MifactError:
        if session.headers.get("Authorization", "").strip():
            return session.headers["Authorization"]
        return login_to_mifact(session)


def mifact_payload(source: dict[str, object], start_date: date, end_date: date) -> dict[str, object]:
    return {
        "monedas": ["PEN", "USD"],
        "lstTipoDocumento": MIFACT_DOCUMENT_TYPES,
        "numeroDocumentoEmisor": source["numero_documento_emisor"],
        "fechaEmisionDesde": f"{start_date.isoformat()}T10:00:00.000Z",
        "fechaEmisionHasta": f"{end_date.isoformat()}T10:00:00.000Z",
        "idEstadoDocumento": 0,
        "dominio": os.getenv("MIFACT_DOMINIO", "sistema.mifact.net"),
        "usuario": (
            os.getenv("MIFACT_USUARIO", "").strip()
            or os.getenv("Mifact_usuario", "").strip()
            or "mrosas.a@outlook.com"
        ),
        "configOrden": {
            "enableFechaEmision": True,
            "fechaEmision": "DESC",
            "enableCorrelativo": True,
            "correlativo": "DESC",
            "enableSerie": True,
            "serie": "ASC",
            "enableTipoCPE": True,
            "tipoCPE": "ASC",
        },
        "numeroSerie": "",
    }


def login_to_xafiro(session: requests.Session) -> None:
    credentials = xafiro_credentials()

    try:
        response = session.post(
            f"{XAFIRO_BASE_URL}/",
            data={**credentials, "enviar": "Ingresar"},
            timeout=REQUEST_TIMEOUT,
        )
        response.raise_for_status()
    except requests.RequestException as exc:
        raise xafiro_request_error(exc, "iniciar sesión") from exc

    # El formulario vuelve a mostrarse cuando Xafiro rechaza las credenciales.
    if re.search(r'name=[\"\']empresa[\"\']', response.text, re.IGNORECASE):
        raise XafiroError("Xafiro rechazó las credenciales configuradas.")


def download_xafiro_report(start_date: date, end_date: date) -> tuple[bytes, str, str]:
    with build_xafiro_session() as session:
        login_to_xafiro(session)

        export_url = (
            f"{XAFIRO_BASE_URL}/ticket/export_ticket_report/"
            f"{quote(start_date.isoformat())}/{quote(end_date.isoformat())}"
        )

        try:
            response = session.get(export_url, timeout=REQUEST_TIMEOUT)
            response.raise_for_status()
        except requests.RequestException as exc:
            raise xafiro_request_error(exc, "descargar el reporte") from exc

        content_type = response.headers.get("Content-Type", "").split(";", 1)[0].strip()
        if "text/html" in content_type.lower():
            raise XafiroError("La sesión de Xafiro terminó antes de descargar el reporte.")
        if not response.content:
            raise XafiroError("Xafiro devolvió un reporte vacío.")

        # Xafiro declara actualmente `application/vnd.ms-excel`, aunque el
        # archivo que entrega es un libro XLSX (contenedor ZIP).
        is_xlsx = response.content.startswith(b"PK\x03\x04")
        extension = ".xlsx" if is_xlsx else ".xls"
        if is_xlsx:
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"facturacion_xafiro_{start_date.isoformat()}_{end_date.isoformat()}{extension}"
        return response.content, filename, content_type or "application/vnd.ms-excel"


def login_to_neo(session: requests.Session) -> None:
    credentials = neo_credentials()

    try:
        response = session.post(
            f"{NEO_AUTH_BASE_URL}/usuario/login",
            json=credentials,
            timeout=REQUEST_TIMEOUT,
        )
        response.raise_for_status()
    except requests.RequestException as exc:
        raise neo_request_error(exc, "iniciar sesión") from exc

    try:
        payload = response.json()
    except ValueError as exc:
        raise NeoError("Neo devolvió una respuesta de login inválida.") from exc

    if not payload.get("estado"):
        raise NeoError(payload.get("message") or "Neo rechazó las credenciales configuradas.")

    token = ((payload.get("data") or {}).get("access_token") or {}).get("token")
    if not token:
        raise NeoError("Neo no devolvió token de acceso para descargar el reporte.")

    session.headers.update({"Authorization": f"Bearer {token}"})


def fetch_neo_report_rows(session: requests.Session, start_date: date, end_date: date) -> list[dict[str, object]]:
    report_url = (
        f"{NEO_BOLETERIA_BASE_URL}/reporte/reporte-detallado-ticket"
        f"?fecha_inicio={quote(start_date.isoformat())}&fecha_fin={quote(end_date.isoformat())}"
    )

    try:
        response = session.get(report_url, timeout=REQUEST_TIMEOUT)
        response.raise_for_status()
    except requests.RequestException as exc:
        raise neo_request_error(exc, "descargar el reporte") from exc

    try:
        payload = response.json()
    except ValueError as exc:
        raise NeoError("Neo devolvió un reporte con formato inválido.") from exc

    if not payload.get("estado"):
        raise NeoError(payload.get("message") or "Neo no pudo generar el reporte solicitado.")

    rows = payload.get("data")
    if not isinstance(rows, list):
        raise NeoError("Neo devolvió el reporte sin filas válidas.")

    return [row for row in rows if isinstance(row, dict)]


def excel_value(value: object) -> object:
    if isinstance(value, (dict, list)):
        import json

        return json.dumps(value, ensure_ascii=False)
    return value


def build_neo_workbook(rows: list[dict[str, object]]) -> bytes:
    columns = [
        "fecha_operacion",
        "hora_operacion",
        "codigo_venta",
        "numero_operacion",
        "id_venta",
        "id_externo",
        "producto",
        "categoria",
        "cantidad",
        "precio",
        "tipo_pago",
        "tipo_caja",
        "tipo_colita",
        "estado_comprobante",
        "comprobantes",
        "detalles_ticket",
        "gratis",
        "validado",
    ]
    extra_columns = sorted({key for row in rows for key in row.keys()} - set(columns))
    columns.extend(extra_columns)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Facturación Neo"
    worksheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="173F33")
    header_font = Font(color="FFFFFF", bold=True)
    worksheet.append([column.replace("_", " ").title() for column in columns])
    worksheet.auto_filter.ref = f"A1:{worksheet.cell(row=1, column=len(columns)).coordinate}"
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in rows:
        worksheet.append([excel_value(row.get(column)) for column in columns])

    for column_cells in worksheet.columns:
        header = str(column_cells[0].value or "")
        max_length = max(len(str(cell.value or "")) for cell in column_cells[:80])
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length, len(header)) + 2,
            42,
        )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def register_neo_export(session: requests.Session, filename: str) -> None:
    try:
        response = session.post(
            f"{NEO_BOLETERIA_BASE_URL}/export-log/trabajar-export-log",
            json={"filename": filename, "origen": "reporte_detallado_tickets"},
            timeout=REQUEST_TIMEOUT,
        )
        response.raise_for_status()
    except requests.RequestException as exc:
        raise neo_request_error(exc, "registrar la exportación") from exc


def download_neo_report(start_date: date, end_date: date) -> tuple[bytes, str, str]:
    filename = f"facturacion_neo_{start_date.isoformat()}_{end_date.isoformat()}.xlsx"

    with build_neo_session() as session:
        login_to_neo(session)
        rows = fetch_neo_report_rows(session, start_date, end_date)
        content = build_neo_workbook(rows)
        register_neo_export(session, filename)

    return (
        content,
        filename,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def culqi_api_request(
    session: requests.Session,
    method: str,
    endpoint: str,
    action: str,
    **kwargs: object,
) -> dict[str, object]:
    try:
        response = session.request(
            method,
            f"{CULQI_API_BASE_URL}/{endpoint.lstrip('/')}",
            timeout=REQUEST_TIMEOUT,
            **kwargs,
        )
    except requests.RequestException as exc:
        raise culqi_request_error(exc, action) from exc

    try:
        payload = response.json()
    except ValueError as exc:
        if response.status_code >= 400:
            raise CulqiError(f"Culqi respondió con error HTTP {response.status_code} al {action}.") from exc
        raise CulqiError(f"Culqi devolvió una respuesta inválida al {action}.") from exc

    if response.status_code >= 400:
        errors = payload.get("errors")
        message = ""
        if isinstance(errors, list) and errors and isinstance(errors[0], dict):
            message = str(errors[0].get("message") or "")
        raise CulqiError(
            f"Culqi respondió con error HTTP {response.status_code} al {action}"
            f"{': ' + message if message else '.'}"
        )

    if payload.get("success") is False:
        raise CulqiError(str(payload.get("message") or f"Culqi no pudo completar: {action}."))

    return payload


def login_to_culqi(session: requests.Session) -> None:
    credentials = culqi_credentials()
    payload = culqi_api_request(
        session,
        "POST",
        "auth/login",
        "iniciar sesión",
        json={**credentials, "rememberMe": False},
    )

    data = payload.get("data")
    if not isinstance(data, dict):
        raise CulqiError("Culqi devolvió una respuesta de login inválida.")

    token = data.get("accessToken")
    session_id = data.get("sessionId")
    if not token or not session_id:
        raise CulqiError("Culqi no devolvió token de acceso para descargar el reporte.")

    session.headers.update(
        {
            "Authorization": f"Bearer {token}",
            "x-session-id": str(session_id),
            "X-CULQI-ENV": "live",
        }
    )


def fetch_culqi_merchants(session: requests.Session) -> list[dict[str, object]]:
    payload = culqi_api_request(session, "GET", "gc/merchants", "obtener comercios")
    data = payload.get("data")
    if not isinstance(data, dict) or not isinstance(data.get("merchants"), list):
        raise CulqiError("Culqi devolvió la lista de comercios con formato inválido.")

    leaves: list[dict[str, object]] = []

    def walk(nodes: list[object], fallback_parent_id: str | None = None) -> None:
        for raw_node in nodes:
            if not isinstance(raw_node, dict):
                continue
            products = raw_node.get("products")
            if isinstance(products, list) and products:
                leaves.append(
                    {
                        **raw_node,
                        "_parent_id": str(raw_node.get("parentId") or fallback_parent_id or raw_node.get("id")),
                    }
                )
            children = raw_node.get("merchants")
            if isinstance(children, list):
                walk(children, str(raw_node.get("id")))

    walk(data["merchants"])
    return leaves


def culqi_products(merchant: dict[str, object], product_name: str) -> list[dict[str, object]]:
    products = merchant.get("products")
    if not isinstance(products, list):
        return []
    return [
        product
        for product in products
        if isinstance(product, dict) and str(product.get("name") or "").strip() == product_name
    ]


def culqi_unix_day(value: date, end_of_day: bool = False) -> int:
    clock = time.max if end_of_day else time.min
    moment = datetime.combine(value, clock, tzinfo=CULQI_TIMEZONE)
    return int(moment.timestamp())


def culqi_iso_day(value: date, end_of_day: bool = False) -> str:
    clock = "23:59:59" if end_of_day else "00:00:00"
    return f"{value.isoformat()}T{clock}Z"


def fetch_culqi_sales_rows(
    session: requests.Session,
    merchant: dict[str, object],
    start_date: date,
    end_date: date,
) -> list[dict[str, object]]:
    products = culqi_products(merchant, "CulqiFull")
    product_ids = [int(product["id"]) for product in products if str(product.get("id") or "").isdigit()]
    merchant_id = str(merchant.get("merchantId") or "")
    parent_id = str(merchant.get("_parent_id") or merchant.get("id") or "")
    if not product_ids or not merchant_id or not parent_id:
        return []

    session.headers["Merchant-Hierarchy-ID"] = parent_id
    rows: list[dict[str, object]] = []
    page = 1
    page_size = 100

    while True:
        body = {
            "startDate": culqi_unix_day(start_date),
            "endDate": culqi_unix_day(end_date, end_of_day=True),
            "merchantIds": [merchant_id],
            "size": page_size,
            "productTypes": product_ids,
            "csiList": ["0", "1"],
            "start": page,
        }
        payload = culqi_api_request(session, "POST", "sales", "descargar ventas CulqiFull", json=body)
        data = payload.get("data")
        if not isinstance(data, dict):
            raise CulqiError("Culqi devolvió ventas con formato inválido.")

        results = data.get("results")
        if not isinstance(results, list):
            raise CulqiError("Culqi devolvió ventas sin filas válidas.")

        for result in results:
            if isinstance(result, dict):
                rows.append(
                    {
                        "origen": "CulqiFull",
                        "comercio_culqi": merchant.get("name"),
                        "merchant_id_culqi": merchant_id,
                        **result,
                    }
                )

        meta = data.get("meta")
        total_records = int((meta or {}).get("total_records") or len(rows)) if isinstance(meta, dict) else len(rows)
        if page * page_size >= total_records or not results:
            break
        page += 1

    return rows


def fetch_culqi_link_rows(
    session: requests.Session,
    merchant: dict[str, object],
    start_date: date,
    end_date: date,
) -> list[dict[str, object]]:
    products = culqi_products(merchant, "CulqiLink")
    public_keys = [str(product["publicKey"]) for product in products if product.get("publicKey")]
    merchant_id = str(merchant.get("merchantId") or "")
    parent_id = str(merchant.get("_parent_id") or merchant.get("id") or "")
    if not public_keys or not merchant_id or not parent_id:
        return []

    session.headers["Merchant-Hierarchy-ID"] = parent_id
    rows: list[dict[str, object]] = []
    page = 1
    page_size = 100

    while True:
        body = {
            "startDate": culqi_iso_day(start_date),
            "endDate": culqi_iso_day(end_date, end_of_day=True),
            "start": page,
            "size": page_size,
            "publicKeys": public_keys,
            "merchantIds": [merchant_id],
        }
        payload = culqi_api_request(session, "POST", "payment-links/list", "descargar links CulqiLink", json=body)
        data = payload.get("data")
        if not isinstance(data, dict):
            raise CulqiError("Culqi devolvió links con formato inválido.")

        links = data.get("links")
        if not isinstance(links, list):
            raise CulqiError("Culqi devolvió links sin filas válidas.")

        for link in links:
            if isinstance(link, dict):
                rows.append(
                    {
                        "origen": "CulqiLink",
                        "comercio_culqi": merchant.get("name"),
                        "merchant_id_culqi": merchant_id,
                        **link,
                    }
                )

        pagination = data.get("pagination")
        total_pages = int((pagination or {}).get("totalPages") or 1) if isinstance(pagination, dict) else 1
        if page >= total_pages or not links:
            break
        page += 1

    return rows


def culqi_raw_request(
    session: requests.Session,
    method: str,
    endpoint: str,
    action: str,
    **kwargs: object,
) -> requests.Response:
    try:
        response = session.request(
            method,
            f"{CULQI_API_BASE_URL}/{endpoint.lstrip('/')}",
            timeout=REQUEST_TIMEOUT,
            **kwargs,
        )
    except requests.RequestException as exc:
        raise culqi_request_error(exc, action) from exc

    if response.status_code >= 400:
        message = ""
        try:
            payload = response.json()
            errors = payload.get("errors") if isinstance(payload, dict) else None
            if isinstance(errors, list) and errors and isinstance(errors[0], dict):
                message = str(errors[0].get("message") or "")
            elif isinstance(payload, dict):
                message = str(payload.get("message") or payload.get("error") or "")
        except ValueError:
            message = response.text[:180].strip()
        raise CulqiError(
            f"Culqi respondió con error HTTP {response.status_code} al {action}"
            f"{': ' + message if message else '.'}"
        )

    return response


def culqi_json_from_response(response: requests.Response, action: str) -> dict[str, object]:
    try:
        payload = response.json()
    except ValueError as exc:
        raise CulqiError(f"Culqi devolvió una respuesta inválida al {action}.") from exc

    if not isinstance(payload, dict):
        raise CulqiError(f"Culqi devolvió una respuesta inválida al {action}.")

    if payload.get("success") is False:
        raise CulqiError(str(payload.get("message") or f"Culqi no pudo completar: {action}."))

    return payload


def find_nested_string(payload: object, keys: set[str]) -> str:
    if isinstance(payload, dict):
        for key, value in payload.items():
            if key in keys and isinstance(value, str) and value.strip():
                return value.strip()
        for value in payload.values():
            found = find_nested_string(value, keys)
            if found:
                return found
    elif isinstance(payload, list):
        for value in payload:
            found = find_nested_string(value, keys)
            if found:
                return found
    return ""


def culqi_report_id(payload: dict[str, object]) -> str:
    return find_nested_string(
        payload,
        {
            "id",
            "reportId",
            "report_id",
            "documentId",
            "document_id",
            "exportDocumentId",
            "export_document_id",
            "downloadId",
            "download_id",
        },
    )


def culqi_report_url(payload: dict[str, object]) -> str:
    return find_nested_string(
        payload,
        {
            "url",
            "fileUrl",
            "file_url",
            "downloadUrl",
            "download_url",
            "reportUrl",
            "report_url",
            "signedUrl",
            "signed_url",
            "presignedUrl",
            "presigned_url",
            "data",
        },
    )


def culqi_report_base64(payload: dict[str, object]) -> str:
    value = find_nested_string(payload, {"base64", "fileBase64", "file_base64", "content"})
    if value.startswith("data:") and "," in value:
        return value.split(",", 1)[1]
    return value


def culqi_workbook_from_payload(
    session: requests.Session,
    payload: dict[str, object],
    source_name: str,
) -> bytes | None:
    raw_base64 = culqi_report_base64(payload)
    if raw_base64:
        try:
            content = base64.b64decode(raw_base64)
        except ValueError:
            content = b""
        if content.startswith(b"PK\x03\x04"):
            return content

    url = culqi_report_url(payload)
    if not url:
        return None

    if url.startswith("/"):
        url = f"{CULQI_API_BASE_URL}{url}"
    elif not url.startswith(("http://", "https://")):
        url = f"{CULQI_API_BASE_URL}/{url.lstrip('/')}"
    try:
        response = session.get(url, timeout=REQUEST_TIMEOUT)
        response.raise_for_status()
    except requests.RequestException as exc:
        raise culqi_request_error(exc, f"descargar el Excel generado de {source_name}") from exc

    if not response.content.startswith(b"PK\x03\x04"):
        raise CulqiError(f"Culqi no devolvió un archivo XLSX válido para {source_name}.")
    return response.content


def download_culqi_generated_workbook(
    session: requests.Session,
    generate_endpoint: str,
    status_endpoint: str,
    body: dict[str, object],
    source_name: str,
    merchant_ids: list[str] | None = None,
) -> bytes:
    response = culqi_raw_request(
        session,
        "POST",
        generate_endpoint,
        f"generar exportable {source_name}",
        json=body,
    )
    if response.content.startswith(b"PK\x03\x04"):
        return response.content

    payload = culqi_json_from_response(response, f"generar exportable {source_name}")
    content = culqi_workbook_from_payload(session, payload, source_name)
    if content:
        return content

    report_id = culqi_report_id(payload)
    if not report_id:
        raise CulqiError(f"Culqi no devolvió el identificador del exportable {source_name}.")

    last_error: CulqiError | None = None
    status_requests: list[tuple[str, str, dict[str, object]]] = []
    if merchant_ids:
        status_requests.extend(
            [
                (
                    "POST",
                    f"{status_endpoint.rstrip('/')}/{quote(report_id)}",
                    {"json": {"merchantIds": merchant_ids}},
                ),
                (
                    "POST",
                    status_endpoint,
                    {"json": {"id": report_id, "merchantIds": merchant_ids}},
                ),
            ]
        )
    status_requests.extend(
        [
            ("GET", f"{status_endpoint.rstrip('/')}/{quote(report_id)}", {}),
            ("GET", status_endpoint, {"params": {"id": report_id}}),
        ]
    )
    for attempt in range(18):
        if attempt:
            time_module.sleep(2)

        for method, endpoint, kwargs in status_requests:
            try:
                status_response = culqi_raw_request(
                    session,
                    method,
                    endpoint,
                    f"consultar exportable {source_name}",
                    **kwargs,
                )
            except CulqiError as exc:
                last_error = exc
                continue

            if status_response.content.startswith(b"PK\x03\x04"):
                return status_response.content
            if not status_response.content:
                continue

            status_payload = culqi_json_from_response(
                status_response,
                f"consultar exportable {source_name}",
            )
            content = culqi_workbook_from_payload(session, status_payload, source_name)
            if content:
                return content

            status_value = str(
                find_nested_string(status_payload, {"status", "state", "estado"}) or ""
            ).lower()
            if status_value in {"failed", "error", "rejected", "fallido"}:
                raise CulqiError(f"Culqi no pudo generar el exportable {source_name}.")

    if last_error:
        raise last_error
    raise CulqiError(f"Culqi tardó demasiado en generar el exportable {source_name}.")


def culqi_sales_report_body(
    merchant_id: str,
    product_ids: list[int],
    start_date: date,
    end_date: date,
) -> dict[str, object]:
    return {
        "typeExport": "full",
        "fileType": "XLSX",
        "docType": "voucher",
        "startDate": culqi_unix_day(start_date),
        "endDate": culqi_unix_day(end_date, end_of_day=True),
        "merchantIds": [merchant_id],
        "size": 100,
        "productTypes": product_ids,
        "statusList": CULQI_SALES_STATUS,
        "merchantStatusList": CULQI_SALES_STATUS,
        "currencyList": CULQI_CURRENCIES,
        "cardTypeList": CULQI_CARD_TYPES,
        "cardBrandList": CULQI_CARD_BRANDS,
        "partialPaymentList": CULQI_PARTIAL_PAYMENTS,
        "walletList": CULQI_WALLETS,
        "csiList": ["0", "1"],
        "start": 1,
        "level": 3,
    }


def download_culqi_full_workbook(
    session: requests.Session,
    merchant: dict[str, object],
    start_date: date,
    end_date: date,
) -> bytes | None:
    products = culqi_products(merchant, "CulqiFull")
    product_ids = [int(product["id"]) for product in products if str(product.get("id") or "").isdigit()]
    merchant_id = str(merchant.get("merchantId") or "")
    parent_id = str(merchant.get("_parent_id") or merchant.get("id") or "")
    if not product_ids or not merchant_id or not parent_id:
        return None

    session.headers["Merchant-Hierarchy-ID"] = parent_id
    body = culqi_sales_report_body(merchant_id, product_ids, start_date, end_date)
    return download_culqi_generated_workbook(
        session,
        "sales/download-report",
        "sales/get-report",
        body,
        "CulqiFull",
        merchant_ids=[merchant_id],
    )


def download_culqi_link_workbook(
    session: requests.Session,
    merchant: dict[str, object],
    start_date: date,
    end_date: date,
) -> bytes | None:
    products = culqi_products(merchant, "CulqiLink")
    product_ids = [int(product["id"]) for product in products if str(product.get("id") or "").isdigit()]
    merchant_id = str(merchant.get("merchantId") or "")
    parent_id = str(merchant.get("_parent_id") or merchant.get("id") or "")
    if not product_ids or not merchant_id or not parent_id:
        return None

    session.headers["Merchant-Hierarchy-ID"] = parent_id
    body = culqi_sales_report_body(merchant_id, product_ids, start_date, end_date)
    return download_culqi_generated_workbook(
        session,
        "sales/download-report",
        "sales/get-report",
        body,
        "CulqiLink",
        merchant_ids=[merchant_id],
    )


def culqi_rows_from_workbook(content: bytes, source_name: str) -> tuple[list[object], list[list[object]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise CulqiError(f"No se pudo leer el Excel de Culqi para {source_name}.") from exc

    worksheet = workbook.active
    rows = [
        list(row)
        for row in worksheet.iter_rows(values_only=True)
        if any(cell is not None and str(cell).strip() for cell in row)
    ]
    if not rows:
        return [], []

    header = trim_trailing_empty_cells(rows[0])
    data_rows = [trim_trailing_empty_cells(row) for row in rows[1:]]
    return header, data_rows


def build_culqi_workbook(source_contents: list[tuple[str, bytes]]) -> bytes:
    header: list[object] | None = None
    merged_rows: list[list[object]] = []

    for source_name, content in source_contents:
        current_header, current_rows = culqi_rows_from_workbook(content, source_name)
        if not current_header:
            continue

        if header is None:
            header = current_header
        elif [str(value or "").strip() for value in header] != [
            str(value or "").strip() for value in current_header
        ]:
            raise CulqiError(
                "Los exportables de CulqiFull y CulqiLink no tienen los mismos encabezados; "
                "no se pueden unir sin mover datos."
            )

        for row in current_rows:
            if len(row) > len(header):
                raise CulqiError(
                    f"El exportable {source_name} tiene más columnas que sus encabezados; "
                    "no se puede unir sin mover datos."
                )
            merged_rows.append(row + [None] * (len(header) - len(row)))

    if header is None:
        header = CULQI_EMPTY_EXPORT_HEADERS

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Facturación Culqi"
    worksheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="173F33")
    header_font = Font(color="FFFFFF", bold=True)
    worksheet.append(header)
    worksheet.auto_filter.ref = f"A1:{worksheet.cell(row=1, column=len(header)).coordinate}"
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in merged_rows:
        worksheet.append(row)

    for column_cells in worksheet.columns:
        header = str(column_cells[0].value or "")
        max_length = max(len(str(cell.value or "")) for cell in column_cells[:80])
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length, len(header)) + 2,
            42,
        )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def download_culqi_report(start_date: date, end_date: date) -> tuple[bytes, str, str]:
    filename = f"facturacion_culqi_{start_date.isoformat()}_{end_date.isoformat()}.xlsx"

    with build_culqi_session() as session:
        login_to_culqi(session)
        merchants = fetch_culqi_merchants(session)
        full_workbooks: list[tuple[str, bytes]] = []
        link_workbooks: list[tuple[str, bytes]] = []
        for merchant in merchants:
            full_content = download_culqi_full_workbook(session, merchant, start_date, end_date)
            if full_content:
                full_workbooks.append(("CulqiFull", full_content))
        for merchant in merchants:
            link_content = download_culqi_link_workbook(session, merchant, start_date, end_date)
            if link_content:
                link_workbooks.append(("CulqiLink", link_content))

    content = build_culqi_workbook(full_workbooks + link_workbooks)
    return (
        content,
        filename,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def download_mifact_source(
    session: requests.Session,
    source: dict[str, object],
    start_date: date,
    end_date: date,
) -> tuple[list[object], list[list[object]]]:
    source_name = str(source.get("name") or "Mifact")
    url = str(source.get("url") or "").strip()
    if not url:
        raise MifactError(f"Falta la URL de Mifact para {source_name}.")

    headers = {"Authorization": resolve_mifact_authorization(session, source)}
    try:
        response = session.post(
            url,
            json=mifact_payload(source, start_date, end_date),
            headers=headers,
            timeout=REQUEST_TIMEOUT,
        )
        response.raise_for_status()
    except requests.RequestException as exc:
        raise mifact_request_error(exc, f"descargar documentos de {source_name}") from exc

    if not response.content:
        raise MifactError(f"Mifact devolvió un reporte vacío para {source_name}.")

    if response.content.startswith(b"PK\x03\x04"):
        return mifact_rows_from_workbook(response.content, source_name)

    content_type = response.headers.get("Content-Type", "").lower()
    if "application/json" in content_type or response.content[:1] in {b"{", b"["}:
        try:
            payload = response.json()
        except ValueError as exc:
            raise MifactError(f"Mifact no devolvió un Excel válido para {source_name}.") from exc
        return mifact_rows_from_json(payload, source_name)

    raise MifactError(f"Mifact no devolvió un archivo XLSX válido para {source_name}.")


def trim_trailing_empty_cells(row: list[object]) -> list[object]:
    while row and (row[-1] is None or str(row[-1]).strip() == ""):
        row.pop()
    return row


CONSOLIDADO_XAFIRO_COLUMNS = {
    "fecha": 8,
    "monto": 18,
    "serie": 3,
    "correlativo": 4,
    "documento": 11,
    "nombre_completo": 12,
}
CONSOLIDADO_XAFIRO_RESORT_AMOUNT_COLUMNS = (22, 24, 25)
CONSOLIDADO_MIFACT_COLUMNS = {
    "fecha": 1,
    "monto": 17,
    "serie": 4,
    "correlativo": 5,
    "documento": 7,
}
CONSOLIDADO_MIFACT_PAYMENT_COLUMN = 25
CONSOLIDADO_MIFACT_EXCLUDED_PAYMENT_METHODS = {"deposito", "depositos", "transferencia", "efectivo"}
CONSOLIDADO_CULQI_COLUMNS = {
    "fecha": 9,
    "monto": 26,
    "nombre": 11,
    "apellido": 12,
}
CONSOLIDADO_CULQI_STATUS_COLUMN = 37
CONSOLIDADO_CULQI_EXCLUDED_STATUSES = {
    "rechazada": "RECHAZADA",
    "anulada": "ANULADA",
}
CONSOLIDADO_EXTRA_HEADERS = [
    "Serie",
    "Correlativo",
    "Documento",
    "Origen comprobante",
    "Estado conciliación",
]
CONSOLIDADO_STATUS_ORDER = {"RECHAZADA": 0, "ANULADA": 0, "REVISAR": 1, "NO ENCONTRADO": 2, "CONCILIADO": 3}
CONSOLIDADO_BANCOS_EXTRA_HEADERS = [
    "Serie",
    "Correlativo",
    "Documento",
    "Estado conciliación",
]
CONSOLIDADO_FACTURADOR_HIGHLIGHT = "F4B183"
CONSOLIDADO_BANCOS_FACTURADOR_HIGHLIGHT = "D9B3FF"
CONSOLIDADO_BANCOS_EXCLUDED_DESCRIPTIONS = {
    "de banco de credito de",
    "de otra cuenta",
    "de joinnus s.a.c",
}
CONSOLIDADO_BANCOS_STATUS_ORDER = {"NO TOMADO": 0, "ANULADO": 1, "NO ENCONTRADO": 2, "ACEPTADO": 3, "CONCILIADO MIFACT": 3}
CONSOLIDADO_TYPES = {"resort", "asociacion", "bancos"}


def normalize_text(value: object) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = re.sub(r"\s+", " ", text)
    return text


def normalize_date_value(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None

    text = str(value).strip()
    if not text:
        return None

    for pattern in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text[:19], pattern).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def normalize_amount_value(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2)

    text = str(value).strip()
    if not text:
        return None
    text = re.sub(r"[^0-9,.-]", "", text)
    if not text:
        return None

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")

    try:
        return round(float(text), 2)
    except ValueError:
        return None


def normalize_digits(value: object) -> str:
    return re.sub(r"\D+", "", str(value or ""))


def operation_matches(left: object, right: object) -> bool:
    left_digits = normalize_digits(left)
    right_digits = normalize_digits(right)
    if not left_digits or not right_digits:
        return False

    if left_digits == right_digits:
        return True

    left_trimmed = left_digits.lstrip("0")
    right_trimmed = right_digits.lstrip("0")
    if left_trimmed and left_trimmed == right_trimmed:
        return True

    return (
        len(left_trimmed) >= 6
        and left_trimmed in right_trimmed
        or len(right_trimmed) >= 6
        and right_trimmed in left_trimmed
    )


def amounts_match(left: object, right: object) -> bool:
    left_amount = normalize_amount_value(left)
    right_amount = normalize_amount_value(right)
    if left_amount is None or right_amount is None:
        return True
    return left_amount == right_amount


def normalized_header(value: object) -> str:
    return normalize_text(value).replace("|", " ")


def find_header_index(rows: list[list[object]], required_terms: tuple[str, ...]) -> int:
    for index, row in enumerate(rows):
        headers = [normalized_header(cell) for cell in row]
        if all(any(term in header for header in headers) for term in required_terms):
            return index
    return 0


def find_column_index(
    header: list[object],
    aliases: tuple[str, ...],
    fallback_index: int,
) -> int:
    normalized_aliases = tuple(normalize_text(alias) for alias in aliases)
    for index, value in enumerate(header):
        text = normalized_header(value)
        if text in normalized_aliases or any(alias in text for alias in normalized_aliases):
            return index
    return fallback_index


def row_cell(row: list[object], zero_based_index: int) -> object:
    return row[zero_based_index] if zero_based_index < len(row) else None


def compact_text(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", normalize_text(value))


def bank_description_matches_transaction(description: object, operation_text: object) -> bool:
    description_compact = compact_text(description)
    if not description_compact:
        return False

    operation = str(operation_text or "")
    bank_text = operation.split("-", 1)[1] if "-" in operation else operation
    bank_compact = compact_text(bank_text)
    return bool(bank_compact and bank_compact in description_compact)


def workbook_rows_from_upload(file_storage: object, source_name: str) -> list[list[object]]:
    filename = str(getattr(file_storage, "filename", "") or "")
    if not filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        raise ConsolidadoError(f"{source_name}: carga un archivo Excel .xlsx válido.")

    try:
        excel_file = getattr(file_storage, "stream", file_storage)
        if hasattr(excel_file, "seek"):
            excel_file.seek(0)
        workbook = load_workbook(excel_file, read_only=True, data_only=True)
    except Exception as exc:
        raise ConsolidadoError(f"{source_name}: no se pudo leer el archivo Excel.") from exc

    worksheet = workbook.active
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    if not any(any(cell is not None and str(cell).strip() for cell in row) for row in rows):
        raise ConsolidadoError(f"{source_name}: el archivo está vacío.")
    return rows


def validate_required_columns(rows: list[list[object]], source_name: str, required_columns: dict[str, int]) -> None:
    max_required = max(required_columns.values())
    max_present = max((len(row) for row in rows), default=0)
    if max_present < max_required:
        raise ConsolidadoError(
            f"{source_name}: faltan columnas requeridas. Se necesita hasta la columna {max_required}."
        )
    data_rows = rows[1:]
    if not any(any(cell is not None and str(cell).strip() for cell in row) for row in data_rows):
        raise ConsolidadoError(f"{source_name}: no tiene registros para procesar.")


def cell_at(row: list[object], one_based_index: int) -> object:
    index = one_based_index - 1
    return row[index] if index < len(row) else None


def parse_invoice_rows(
    rows: list[list[object]],
    source_name: str,
    columns: dict[str, int],
    include_name: bool = False,
    excluded_values_by_column: dict[int, set[str]] | None = None,
    amount_sum_columns: tuple[int, ...] | None = None,
) -> list[dict[str, object]]:
    required_columns = dict(columns)
    if amount_sum_columns:
        required_columns.update({f"monto_suma_{index}": column for index, column in enumerate(amount_sum_columns)})
    validate_required_columns(rows, source_name, required_columns)
    invoices: list[dict[str, object]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        if excluded_values_by_column and any(
            normalize_text(cell_at(row, column)) in excluded_values
            for column, excluded_values in excluded_values_by_column.items()
        ):
            continue
        normalized_date = normalize_date_value(cell_at(row, columns["fecha"]))
        if amount_sum_columns:
            amount_parts = [normalize_amount_value(cell_at(row, column)) for column in amount_sum_columns]
            normalized_amount = round(sum(part or 0 for part in amount_parts), 2) if any(part is not None for part in amount_parts) else None
        else:
            normalized_amount = normalize_amount_value(cell_at(row, columns["monto"]))
        if normalized_date is None or normalized_amount is None:
            continue
        invoice = {
            "row_number": row_number,
            "date": normalized_date,
            "amount": normalized_amount,
            "serie": cell_at(row, columns["serie"]),
            "correlativo": cell_at(row, columns["correlativo"]),
            "documento": cell_at(row, columns["documento"]),
        }
        if include_name:
            invoice["name"] = normalize_text(cell_at(row, columns["nombre_completo"]))
        invoices.append(invoice)
    if not invoices:
        raise ConsolidadoError(f"{source_name}: no se encontraron filas con fecha y monto válidos.")
    return invoices


def parse_culqi_rows(rows: list[list[object]], source_name: str) -> tuple[list[object], list[dict[str, object]]]:
    validate_required_columns(rows, source_name, CONSOLIDADO_CULQI_COLUMNS)
    header = trim_trailing_empty_cells(list(rows[0]))
    records: list[dict[str, object]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        row_values = list(row[: len(header)]) + [None] * max(len(header) - len(row), 0)
        culqi_status = CONSOLIDADO_CULQI_EXCLUDED_STATUSES.get(
            normalize_text(cell_at(row, CONSOLIDADO_CULQI_STATUS_COLUMN))
        )
        records.append(
            {
                "source": source_name,
                "row_number": row_number,
                "date": normalize_date_value(cell_at(row, CONSOLIDADO_CULQI_COLUMNS["fecha"])),
                "amount": normalize_amount_value(cell_at(row, CONSOLIDADO_CULQI_COLUMNS["monto"])),
                "name": normalize_text(
                    f"{cell_at(row, CONSOLIDADO_CULQI_COLUMNS['apellido']) or ''} "
                    f"{cell_at(row, CONSOLIDADO_CULQI_COLUMNS['nombre']) or ''}"
                ),
                "row": row_values,
                "serie": "",
                "correlativo": "",
                "documento": "",
                "origen": "",
                "estado": culqi_status or "NO ENCONTRADO",
                "excluded_from_consolidation": culqi_status is not None,
            }
        )
    if not records:
        raise ConsolidadoError(f"{source_name}: no tiene operaciones para consolidar.")
    return header, records


def candidate_records(culqi_records: list[dict[str, object]], invoice: dict[str, object]) -> list[dict[str, object]]:
    return [
        record
        for record in culqi_records
        if record["estado"] == "NO ENCONTRADO"
        and record["date"] == invoice["date"]
        and record["amount"] == invoice["amount"]
    ]


def invoice_match_key(invoice: dict[str, object]) -> tuple[object, object]:
    return invoice["date"], invoice["amount"]


def group_invoices_by_match_key(invoices: list[dict[str, object]]) -> list[list[dict[str, object]]]:
    grouped: dict[tuple[object, object], list[dict[str, object]]] = {}
    for invoice in invoices:
        grouped.setdefault(invoice_match_key(invoice), []).append(invoice)
    return list(grouped.values())


def apply_balanced_duplicate_matches(
    candidates: list[dict[str, object]],
    invoices: list[dict[str, object]],
    source_name: str,
) -> int:
    candidates_by_order = sorted(candidates, key=lambda record: int(record.get("sequence") or 0))
    invoices_by_order = sorted(invoices, key=lambda invoice: int(invoice.get("row_number") or 0))[: len(candidates_by_order)]
    for record, invoice in zip(candidates_by_order, invoices_by_order):
        apply_invoice_match(record, invoice, source_name)
    return len(invoices_by_order)


def name_match_score(invoice_name: str, culqi_name: str) -> int:
    if not invoice_name or not culqi_name:
        return 0
    if invoice_name == culqi_name:
        return 4
    invoice_tokens = set(invoice_name.split())
    culqi_tokens = set(culqi_name.split())
    if not invoice_tokens or not culqi_tokens:
        return 0
    overlap = invoice_tokens & culqi_tokens
    if invoice_tokens <= culqi_tokens or culqi_tokens <= invoice_tokens:
        return 3
    if len(overlap) >= min(3, len(invoice_tokens), len(culqi_tokens)):
        return 2
    return 0


def mark_review(records: list[dict[str, object]]) -> None:
    for record in records:
        if record["estado"] == "NO ENCONTRADO":
            record["estado"] = "REVISAR"


def apply_invoice_match(record: dict[str, object], invoice: dict[str, object], source_name: str) -> None:
    record["serie"] = invoice["serie"]
    record["correlativo"] = invoice["correlativo"]
    record["documento"] = invoice["documento"]
    record["origen"] = source_name
    record["invoice_row_number"] = invoice.get("row_number")
    record["estado"] = "CONCILIADO"


def remaining_invoice_group(
    invoice_groups: dict[tuple[object, object], list[dict[str, object]]],
    invoice: dict[str, object],
    matched_invoice_rows: set[int],
) -> list[dict[str, object]]:
    return [
        peer_invoice
        for peer_invoice in invoice_groups[invoice_match_key(invoice)]
        if int(peer_invoice.get("row_number") or 0) not in matched_invoice_rows
    ]


def mark_invoices_matched(invoices: list[dict[str, object]], matched_invoice_rows: set[int]) -> None:
    for invoice in invoices:
        matched_invoice_rows.add(int(invoice.get("row_number") or 0))


def reconcile_xafiro(culqi_records: list[dict[str, object]], invoices: list[dict[str, object]]) -> int:
    matched = 0
    invoice_groups = {invoice_match_key(group[0]): group for group in group_invoices_by_match_key(invoices)}
    matched_invoice_rows: set[int] = set()
    for invoice in invoices:
        invoice_row = int(invoice.get("row_number") or 0)
        if invoice_row in matched_invoice_rows:
            continue
        candidates = candidate_records(culqi_records, invoice)
        if len(candidates) == 1:
            apply_invoice_match(candidates[0], invoice, "Xafiro")
            matched_invoice_rows.add(invoice_row)
            matched += 1
            continue
        if len(candidates) <= 1:
            continue

        invoice_group = remaining_invoice_group(invoice_groups, invoice, matched_invoice_rows)
        if len(candidates) <= len(invoice_group):
            matched += apply_balanced_duplicate_matches(candidates, invoice_group, "Xafiro")
            mark_invoices_matched(invoice_group[: len(candidates)], matched_invoice_rows)
            continue

        scored = [(name_match_score(str(invoice.get("name") or ""), str(record.get("name") or "")), record) for record in candidates]
        best_score = max(score for score, _record in scored)
        best_records = [record for score, record in scored if score == best_score and score > 0]
        if len(best_records) == 1:
            apply_invoice_match(best_records[0], invoice, "Xafiro")
            matched_invoice_rows.add(invoice_row)
            matched += 1
        else:
            mark_review(candidates)
    return matched


def reconcile_mifact(culqi_records: list[dict[str, object]], invoices: list[dict[str, object]]) -> int:
    matched = 0
    invoice_groups = {invoice_match_key(group[0]): group for group in group_invoices_by_match_key(invoices)}
    matched_invoice_rows: set[int] = set()
    for invoice in invoices:
        invoice_row = int(invoice.get("row_number") or 0)
        if invoice_row in matched_invoice_rows:
            continue
        candidates = candidate_records(culqi_records, invoice)
        if len(candidates) == 1:
            apply_invoice_match(candidates[0], invoice, "Mifact")
            matched_invoice_rows.add(invoice_row)
            matched += 1
        elif len(candidates) > 1:
            invoice_group = remaining_invoice_group(invoice_groups, invoice, matched_invoice_rows)
            matched_count = apply_balanced_duplicate_matches(candidates, invoice_group, "Mifact")
            matched += matched_count
            mark_invoices_matched(invoice_group[:matched_count], matched_invoice_rows)
    return matched


def matched_invoice_rows_by_source(records: list[dict[str, object]], source_name: str) -> set[int]:
    return {
        int(record.get("invoice_row_number") or 0)
        for record in records
        if record.get("origen") == source_name and record.get("estado") == "CONCILIADO" and record.get("invoice_row_number")
    }


def append_facturador_sheet(
    workbook: Workbook,
    title: str,
    rows: list[list[object]],
    highlighted_rows: set[int],
    highlight_color: str,
) -> None:
    worksheet = workbook.create_sheet(title[:31])
    fill = PatternFill("solid", fgColor=highlight_color)
    for row_number, row in enumerate(rows, start=1):
        worksheet.append(list(row))
        if row_number in highlighted_rows:
            for cell in worksheet[worksheet.max_row]:
                cell.fill = fill

    if rows:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{worksheet.cell(row=1, column=max(len(rows[0]), 1)).coordinate}"
    for column_cells in worksheet.columns:
        header_value = str(column_cells[0].value or "")
        max_length = max(len(str(cell.value or "")) for cell in column_cells[:80])
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length, len(header_value)) + 2, 42)


def build_consolidado_workbook(
    header: list[object],
    records: list[dict[str, object]],
    facturador_sheets: list[tuple[str, list[list[object]], set[int]]] | None = None,
    highlight_color: str = CONSOLIDADO_FACTURADOR_HIGHLIGHT,
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Consolidado Culqi"
    worksheet.freeze_panes = "A2"

    final_header = header + CONSOLIDADO_EXTRA_HEADERS
    worksheet.append(final_header)
    worksheet.auto_filter.ref = f"A1:{worksheet.cell(row=1, column=len(final_header)).coordinate}"

    header_fill = PatternFill("solid", fgColor="173F33")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    sorted_records = sorted(
        records,
        key=lambda record: (CONSOLIDADO_STATUS_ORDER.get(str(record["estado"]), 9), int(record.get("sequence", record["row_number"]))),
    )
    for record in sorted_records:
        worksheet.append(
            list(record["row"])
            + [record["serie"], record["correlativo"], record["documento"], record["origen"], record["estado"]]
        )
        if record.get("excluded_from_consolidation"):
            for cell in worksheet[worksheet.max_row]:
                cell.font = Font(color="FFFF0000")

    for column_cells in worksheet.columns:
        header_value = str(column_cells[0].value or "")
        max_length = max(len(str(cell.value or "")) for cell in column_cells[:80])
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length, len(header_value)) + 2, 42)

    for title, rows, highlighted_rows in facturador_sheets or []:
        append_facturador_sheet(workbook, title, rows, highlighted_rows, highlight_color)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def parse_bancos_rows(rows: list[list[object]]) -> tuple[list[list[object]], list[object], list[dict[str, object]]]:
    header_index = find_header_index(rows, ("fecha", "monto", "operacion"))
    header = trim_trailing_empty_cells(list(rows[header_index]))
    date_index = find_column_index(header, ("fecha",), 0)
    description_index = find_column_index(header, ("descripcion operacion", "descripci n operaci n", "operacion"), 2)
    amount_index = find_column_index(header, ("monto",), 6)
    operation_index = find_column_index(header, ("operacion - numero", "operacion numero", "numero"), 9)

    records: list[dict[str, object]] = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue

        description = row_cell(row, description_index)
        amount = normalize_amount_value(row_cell(row, amount_index))
        row_values = list(row[: len(header)]) + [None] * max(len(header) - len(row), 0)
        no_tomado = amount is not None and amount < 0 or normalize_text(description) in CONSOLIDADO_BANCOS_EXCLUDED_DESCRIPTIONS
        records.append(
            {
                "row_number": row_number,
                "row": row_values,
                "date": normalize_date_value(row_cell(row, date_index)),
                "description": description,
                "operation": row_cell(row, operation_index),
                "amount": amount,
                "serie": "",
                "correlativo": "",
                "documento": "",
                "estado": "NO TOMADO" if no_tomado else "NO ENCONTRADO",
                "needs_review": True,
                "excluded_from_conciliation": no_tomado,
            }
        )

    if not records:
        raise ConsolidadoError("Bancos: no se encontraron operaciones positivas para consolidar.")
    return rows[:header_index], header, records


def parse_xafiro_transfer_rows(rows: list[list[object]]) -> list[dict[str, object]]:
    header_index = find_header_index(rows, ("reserva", "medio de pago", "operacion"))
    header = list(rows[header_index])
    date_index = find_column_index(header, ("fecha",), 0)
    reserve_index = find_column_index(header, ("n reserva", "reserva"), 1)
    amount_index = find_column_index(header, ("monto",), 3)
    payment_index = find_column_index(header, ("medio de pago",), 5)
    voucher_index = find_column_index(header, ("voucher operacion", "operacion"), 6)

    transfers: list[dict[str, object]] = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        if normalize_text(row_cell(row, payment_index)) != "transferencia":
            continue
        transfers.append(
            {
                "row_number": row_number,
                "date": normalize_date_value(row_cell(row, date_index)),
                "reserve": row_cell(row, reserve_index),
                "amount": normalize_amount_value(row_cell(row, amount_index)),
                "operation": row_cell(row, voucher_index),
            }
        )
    return transfers


def parse_xafiro_facturacion_rows(rows: list[list[object]]) -> list[dict[str, object]]:
    header_index = find_header_index(rows, ("serie", "documento", "estado"))
    header = list(rows[header_index])
    serie_index = find_column_index(header, ("serie",), 2)
    invoice_number_index = find_column_index(header, ("numero",), 3)
    reserve_index = find_column_index(header, ("codigo",), 5)
    transfer_amount_index = find_column_index(header, ("transferencia",), 22)
    document_index = next(
        (index for index, value in enumerate(header) if normalized_header(value) == "documento"),
        10,
    )
    status_index = find_column_index(header, ("estado",), 27)

    invoices: list[dict[str, object]] = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        reserve_digits = normalize_digits(row_cell(row, reserve_index)).lstrip("0")
        if not reserve_digits:
            continue
        invoices.append(
            {
                "row_number": row_number,
                "reserve": reserve_digits,
                "serie": row_cell(row, serie_index),
                "correlativo": row_cell(row, invoice_number_index),
                "documento": row_cell(row, document_index),
                "transfer_amount": normalize_amount_value(row_cell(row, transfer_amount_index)),
                "estado": str(row_cell(row, status_index) or "").strip() or "SIN ESTADO",
            }
        )
    return invoices


def parse_mifact_transfer_rows_for_bancos(rows: list[list[object]]) -> list[dict[str, object]]:
    header_index = find_header_index(rows, ("serie", "correlativo", "forma de pago"))
    header = list(rows[header_index])
    status_index = find_column_index(header, ("estado documento",), 2)
    serie_index = find_column_index(header, ("serie",), 3)
    correlativo_index = find_column_index(header, ("correlativo",), 4)
    document_index = find_column_index(header, ("ruc / dni: cliente", "cliente"), 6)
    amount_index = find_column_index(header, ("venta total",), 16)
    payment_index = find_column_index(header, ("forma de pago",), 24)
    operation_index = find_column_index(header, ("observacion",), 29)

    invoices: list[dict[str, object]] = []
    for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        payment = normalize_text(row_cell(row, payment_index))
        if payment not in {"transferencia", "deposito", "depositos"}:
            continue
        invoices.append(
            {
                "row_number": row_number,
                "operation": row_cell(row, operation_index),
                "amount": normalize_amount_value(row_cell(row, amount_index)),
                "serie": row_cell(row, serie_index),
                "correlativo": row_cell(row, correlativo_index),
                "documento": row_cell(row, document_index),
                "estado": str(row_cell(row, status_index) or "").strip() or "CONCILIADO",
            }
        )
    return invoices


def choose_facturacion_invoice(
    reserve: object,
    bank_amount: object,
    invoices: list[dict[str, object]],
    matched_invoice_rows: set[int],
) -> dict[str, object] | None:
    reserve_key = normalize_digits(reserve).lstrip("0")
    if not reserve_key:
        return None

    candidates = [
        invoice
        for invoice in invoices
        if invoice["reserve"] == reserve_key and int(invoice.get("row_number") or 0) not in matched_invoice_rows
        and normalize_amount_value(invoice.get("transfer_amount")) == normalize_amount_value(bank_amount)
    ]
    if not candidates:
        return None

    accepted = [invoice for invoice in candidates if normalize_text(invoice.get("estado")) == "aceptado"]
    return accepted[0] if accepted else candidates[0]


def apply_bancos_invoice(
    record: dict[str, object],
    invoice: dict[str, object],
    status: str,
    needs_review: bool,
    source_name: str,
    transaction_row_number: object | None = None,
) -> None:
    record["serie"] = invoice["serie"]
    record["correlativo"] = invoice["correlativo"]
    record["documento"] = invoice["documento"]
    record["estado"] = status
    record["needs_review"] = needs_review
    record["facturador_origen"] = source_name
    record["invoice_row_number"] = invoice.get("row_number")
    if transaction_row_number is not None:
        record["transaction_row_number"] = transaction_row_number


def reconcile_bancos_xafiro(
    records: list[dict[str, object]],
    transfers: list[dict[str, object]],
    invoices: list[dict[str, object]],
) -> int:
    matched = 0
    matched_transfer_rows: set[int] = set()
    matched_invoice_rows: set[int] = set()

    for record in records:
        if record["estado"] != "NO ENCONTRADO":
            continue
        transfer = next(
            (
                transfer
                for transfer in transfers
                if int(transfer.get("row_number") or 0) not in matched_transfer_rows
                and operation_matches(record.get("operation"), transfer.get("operation"))
                and amounts_match(record.get("amount"), transfer.get("amount"))
            ),
            None,
        )
        if transfer is None:
            continue

        invoice = choose_facturacion_invoice(transfer.get("reserve"), record.get("amount"), invoices, matched_invoice_rows)
        if invoice is None:
            continue

        matched_transfer_rows.add(int(transfer.get("row_number") or 0))
        matched_invoice_rows.add(int(invoice.get("row_number") or 0))
        invoice_status = str(invoice.get("estado") or "").strip() or "SIN ESTADO"
        if normalize_text(invoice_status) == "aceptado":
            apply_bancos_invoice(record, invoice, "ACEPTADO", False, "Xafiro", transfer.get("row_number"))
            matched += 1
        else:
            apply_bancos_invoice(record, invoice, invoice_status.upper(), True, "Xafiro", transfer.get("row_number"))

    for record in records:
        if record["estado"] != "NO ENCONTRADO":
            continue
        transfer = next(
            (
                transfer
                for transfer in transfers
                if int(transfer.get("row_number") or 0) not in matched_transfer_rows
                and record.get("date") == transfer.get("date")
                and amounts_match(record.get("amount"), transfer.get("amount"))
                and bank_description_matches_transaction(record.get("description"), transfer.get("operation"))
            ),
            None,
        )
        if transfer is None:
            continue

        invoice = choose_facturacion_invoice(transfer.get("reserve"), record.get("amount"), invoices, matched_invoice_rows)
        if invoice is None:
            continue

        matched_transfer_rows.add(int(transfer.get("row_number") or 0))
        matched_invoice_rows.add(int(invoice.get("row_number") or 0))
        invoice_status = str(invoice.get("estado") or "").strip() or "SIN ESTADO"
        if normalize_text(invoice_status) == "aceptado":
            apply_bancos_invoice(record, invoice, "ACEPTADO", False, "Xafiro", transfer.get("row_number"))
            matched += 1
        else:
            apply_bancos_invoice(record, invoice, invoice_status.upper(), True, "Xafiro", transfer.get("row_number"))

    return matched


def reconcile_bancos_mifact(records: list[dict[str, object]], invoices: list[dict[str, object]]) -> int:
    matched = 0
    matched_invoice_rows: set[int] = set()
    for record in records:
        if record["estado"] != "NO ENCONTRADO":
            continue
        invoice = next(
            (
                candidate
                for candidate in invoices
                if int(candidate.get("row_number") or 0) not in matched_invoice_rows
                and operation_matches(record.get("operation"), candidate.get("operation"))
                and amounts_match(record.get("amount"), candidate.get("amount"))
            ),
            None,
        )
        if invoice is None:
            continue
        matched_invoice_rows.add(int(invoice.get("row_number") or 0))
        apply_bancos_invoice(record, invoice, "CONCILIADO MIFACT", False, "Mifact")
        matched += 1
    return matched


def matched_bancos_invoice_rows(records: list[dict[str, object]], source_name: str) -> set[int]:
    return {
        int(record.get("invoice_row_number") or 0)
        for record in records
        if record.get("facturador_origen") == source_name
        and record.get("invoice_row_number")
        and (record.get("estado") == "ACEPTADO" or record.get("estado") == "CONCILIADO MIFACT")
    }


def matched_bancos_transaction_rows(records: list[dict[str, object]]) -> set[int]:
    return {
        int(record.get("transaction_row_number") or 0)
        for record in records
        if record.get("facturador_origen") == "Xafiro"
        and record.get("transaction_row_number")
        and record.get("estado") == "ACEPTADO"
    }


def build_bancos_workbook(
    metadata_rows: list[list[object]],
    header: list[object],
    records: list[dict[str, object]],
    facturador_sheets: list[tuple[str, list[list[object]], set[int]]] | None = None,
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Consolidado Bancos"

    for row in metadata_rows:
        worksheet.append(list(row[: len(header)]) + [None] * max(len(header) - len(row), 0))

    final_header = header + CONSOLIDADO_BANCOS_EXTRA_HEADERS
    worksheet.append(final_header)
    header_row = worksheet.max_row
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.auto_filter.ref = f"A{header_row}:{worksheet.cell(row=header_row, column=len(final_header)).coordinate}"

    header_fill = PatternFill("solid", fgColor="173F33")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[header_row]:
        cell.fill = header_fill
        cell.font = header_font

    sorted_records = sorted(
        records,
        key=lambda record: (
            CONSOLIDADO_BANCOS_STATUS_ORDER.get(str(record.get("estado")), 9),
            int(record.get("row_number") or 0),
        ),
    )
    for record in sorted_records:
        worksheet.append(
            list(record["row"])
            + [record["serie"], record["correlativo"], record["documento"], record["estado"]]
        )
        if record.get("needs_review"):
            for cell in worksheet[worksheet.max_row]:
                cell.font = Font(color="FFFF0000")

    for column_cells in worksheet.columns:
        header_value = str(column_cells[header_row - 1].value or "")
        max_length = max(len(str(cell.value or "")) for cell in column_cells[:80])
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length, len(header_value)) + 2, 42)

    for title, rows, highlighted_rows in facturador_sheets or []:
        append_facturador_sheet(workbook, title, rows, highlighted_rows, CONSOLIDADO_BANCOS_FACTURADOR_HIGHLIGHT)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def process_bancos_consolidado_uploads(files: dict[str, object]) -> tuple[bytes, dict[str, int]]:
    required = ("bancos", "xafiro_transacciones", "xafiro", "mifact")
    missing = [name for name in required if name not in files or not getattr(files.get(name), "filename", "")]
    if missing:
        raise ConsolidadoError("Carga los archivos requeridos: Oficina, Xafiro Transacciones, Xafiro Facturación y Mifact.")

    bancos_rows = workbook_rows_from_upload(files["bancos"], "Oficina")
    xafiro_transfer_rows = workbook_rows_from_upload(files["xafiro_transacciones"], "Xafiro Transacciones")
    xafiro_facturacion_rows = workbook_rows_from_upload(files["xafiro"], "Xafiro Facturación")
    mifact_rows = workbook_rows_from_upload(files["mifact"], "Mifact")

    metadata_rows, bancos_header, bancos_records = parse_bancos_rows(bancos_rows)
    xafiro_transfers = parse_xafiro_transfer_rows(xafiro_transfer_rows)
    xafiro_invoices = parse_xafiro_facturacion_rows(xafiro_facturacion_rows)
    mifact_invoices = parse_mifact_transfer_rows_for_bancos(mifact_rows)

    xafiro_count = reconcile_bancos_xafiro(bancos_records, xafiro_transfers, xafiro_invoices)
    mifact_count = reconcile_bancos_mifact(bancos_records, mifact_invoices)
    no_encontrado = sum(1 for record in bancos_records if record["estado"] == "NO ENCONTRADO")
    revisar = sum(1 for record in bancos_records if record.get("needs_review") and record["estado"] not in {"NO ENCONTRADO", "NO TOMADO"})
    total_conciliado = xafiro_count + mifact_count
    summary = {
        "total_operaciones_culqi": len(bancos_records),
        "conciliados_xafiro": xafiro_count,
        "conciliados_mifact": mifact_count,
        "no_encontrados": no_encontrado,
        "registros_revisar": revisar,
        "total_conciliado": total_conciliado,
        "total_no_conciliado": no_encontrado + revisar,
    }
    facturador_sheets = [
        ("Xafiro Transacciones", xafiro_transfer_rows, matched_bancos_transaction_rows(bancos_records)),
        ("Xafiro Facturacion", xafiro_facturacion_rows, matched_bancos_invoice_rows(bancos_records, "Xafiro")),
        ("Mifact", mifact_rows, matched_bancos_invoice_rows(bancos_records, "Mifact")),
    ]
    return build_bancos_workbook(metadata_rows, bancos_header, bancos_records, facturador_sheets), summary


def process_consolidado_uploads(
    files: dict[str, object],
    consolidation_type: str = "resort",
) -> tuple[bytes, dict[str, int]]:
    consolidation_type = normalize_text(consolidation_type)
    if consolidation_type not in CONSOLIDADO_TYPES:
        raise ConsolidadoError("Selecciona un tipo de consolidado válido: Resort, Asociación o Bancos.")

    if consolidation_type == "bancos":
        return process_bancos_consolidado_uploads(files)

    is_resort = consolidation_type == "resort"
    required = ("xafiro", "mifact", "culqilink", "culqifull") if is_resort else ("mifact", "culqilink", "culqifull")
    missing = [name for name in required if name not in files or not getattr(files.get(name), "filename", "")]
    if missing:
        required_labels = "Xafiro, Mifact, CulqiLink y CulqiFull" if is_resort else "Mifact, CulqiLink y CulqiFull"
        raise ConsolidadoError(f"Carga los archivos requeridos: {required_labels}.")

    mifact_rows = workbook_rows_from_upload(files["mifact"], "Mifact")
    culqilink_rows = workbook_rows_from_upload(files["culqilink"], "CulqiLink")
    culqifull_rows = workbook_rows_from_upload(files["culqifull"], "CulqiFull")

    culqilink_header, culqilink_records = parse_culqi_rows(culqilink_rows, "CulqiLink")
    culqifull_header, culqifull_records = parse_culqi_rows(culqifull_rows, "CulqiFull")
    if [str(value or "").strip() for value in culqilink_header] != [str(value or "").strip() for value in culqifull_header]:
        raise ConsolidadoError("CulqiLink y CulqiFull no tienen los mismos encabezados.")

    culqi_records = culqilink_records + culqifull_records
    for sequence, record in enumerate(culqi_records):
        record["sequence"] = sequence
    mifact_invoices = parse_invoice_rows(
        mifact_rows,
        "Mifact",
        CONSOLIDADO_MIFACT_COLUMNS,
        excluded_values_by_column={CONSOLIDADO_MIFACT_PAYMENT_COLUMN: CONSOLIDADO_MIFACT_EXCLUDED_PAYMENT_METHODS},
    )

    xafiro_count = 0
    xafiro_rows: list[list[object]] | None = None
    if is_resort:
        xafiro_rows = workbook_rows_from_upload(files["xafiro"], "Xafiro")
        xafiro_invoices = parse_invoice_rows(
            xafiro_rows,
            "Xafiro",
            CONSOLIDADO_XAFIRO_COLUMNS,
            include_name=True,
            amount_sum_columns=CONSOLIDADO_XAFIRO_RESORT_AMOUNT_COLUMNS,
        )
        xafiro_count = reconcile_xafiro(culqi_records, xafiro_invoices)
    mifact_count = reconcile_mifact(culqi_records, mifact_invoices)
    no_encontrado = sum(1 for record in culqi_records if record["estado"] == "NO ENCONTRADO")
    revisar = sum(1 for record in culqi_records if record["estado"] == "REVISAR")
    total_conciliado = xafiro_count + mifact_count
    summary = {
        "total_operaciones_culqi": len(culqi_records),
        "conciliados_xafiro": xafiro_count,
        "conciliados_mifact": mifact_count,
        "no_encontrados": no_encontrado,
        "registros_revisar": revisar,
        "total_conciliado": total_conciliado,
        "total_no_conciliado": no_encontrado + revisar,
    }
    facturador_sheets: list[tuple[str, list[list[object]], set[int]]] = []
    if xafiro_rows is not None:
        facturador_sheets.append(("Xafiro", xafiro_rows, matched_invoice_rows_by_source(culqi_records, "Xafiro")))
    facturador_sheets.append(("Mifact", mifact_rows, matched_invoice_rows_by_source(culqi_records, "Mifact")))
    return build_consolidado_workbook(culqilink_header, culqi_records, facturador_sheets), summary


def mifact_rows_from_workbook(content: bytes, source_name: str) -> tuple[list[object], list[list[object]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise MifactError(f"No se pudo leer el Excel de Mifact para {source_name}.") from exc

    worksheet = workbook.active
    rows = [
        list(row)
        for row in worksheet.iter_rows(values_only=True)
        if any(cell is not None and str(cell).strip() for cell in row)
    ]
    if not rows:
        return [], []

    header = trim_trailing_empty_cells(rows[0])
    data_rows = [trim_trailing_empty_cells(row) for row in rows[1:]]
    return header, data_rows


def mifact_rows_from_json(payload: object, source_name: str) -> tuple[list[object], list[list[object]]]:
    if isinstance(payload, dict):
        message = payload.get("message") or payload.get("mensaje") or payload.get("error")
        for value in payload.values():
            if isinstance(value, list):
                payload = value
                break
        else:
            raise MifactError(str(message or f"Mifact no pudo exportar documentos de {source_name}."))

    if not isinstance(payload, list):
        raise MifactError(f"Mifact devolvió una respuesta inválida para {source_name}.")

    records = [row for row in payload if isinstance(row, dict)]
    columns: list[object] = []
    for record in records:
        for key in record.keys():
            if key not in columns:
                columns.append(key)

    if not columns:
        return [], []

    rows = [[excel_value(record.get(str(column))) for column in columns] for record in records]
    return columns, rows


def build_mifact_workbook(source_contents: list[tuple[str, object]]) -> bytes:
    header: list[object] | None = None
    merged_rows: list[list[object]] = []

    for source_name, content in source_contents:
        if isinstance(content, bytes):
            current_header, current_rows = mifact_rows_from_workbook(content, source_name)
        else:
            current_header, current_rows = content
        if not current_header:
            continue

        if header is None:
            header = current_header
        elif [str(value or "").strip() for value in header] != [
            str(value or "").strip() for value in current_header
        ]:
            raise MifactError(
                "Los exportables de Mifact no tienen los mismos campos, por eso no se pueden unir sin cambiar columnas."
            )

        merged_rows.extend(current_rows)

    if header is None:
        raise MifactError("Mifact no devolvió campos para unir el exportable.")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Facturación Mifact"
    worksheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="173F33")
    header_font = Font(color="FFFFFF", bold=True)
    worksheet.append(header)
    worksheet.auto_filter.ref = f"A1:{worksheet.cell(row=1, column=len(header)).coordinate}"
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in merged_rows:
        normalized_row = row[: len(header)] + [None] * max(len(header) - len(row), 0)
        worksheet.append(normalized_row)

    for column_cells in worksheet.columns:
        header_value = str(column_cells[0].value or "")
        max_length = max(len(str(cell.value or "")) for cell in column_cells[:80])
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length, len(header_value)) + 2,
            42,
        )

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def download_mifact_report(start_date: date, end_date: date) -> tuple[bytes, str, str]:
    filename = f"facturacion_mifact_{start_date.isoformat()}_{end_date.isoformat()}.xlsx"

    with build_mifact_session() as session:
        source_contents = [
            (str(source["name"]), download_mifact_source(session, source, start_date, end_date))
            for source in MIFACT_SOURCES
        ]

    content = build_mifact_workbook(source_contents)
    return (
        content,
        filename,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/")
@permission_required("reportes.ver")
def index():
    return render_template("index.html")


@app.get("/animaciones/<path:filename>")
def animation_file(filename: str):
    return send_from_directory(BASE_DIR / "animaciones", filename, max_age=0)


@app.get("/api/health")
def health():
    credentials_ready = all(
        os.getenv(key, "").strip()
        for key in ("XAFIRO_EMPRESA", "XAFIRO_USUARIO", "XAFIRO_PASSWORD")
    )
    neo_ready = bool(
        (os.getenv("NEO_USUARIO", "").strip() or os.getenv("Neo_usuario", "").strip())
        and (os.getenv("NEO_PASSWORD", "").strip() or os.getenv("Neo_password", "").strip())
    )
    culqi_ready = bool(
        (
            os.getenv("CULQI_USUARIO", "").strip()
            or os.getenv("CULQUI_USUARIO", "").strip()
            or os.getenv("Culqui_usuario", "").strip()
        )
        and (
            os.getenv("CULQI_PASSWORD", "").strip()
            or os.getenv("CULQUI_PASSWORD", "").strip()
            or os.getenv("Culqui_password", "").strip()
        )
    )
    mifact_ready = all(
        any(os.getenv(str(env_name), "").strip() for env_name in source.get("token_envs", ()))
        for source in MIFACT_SOURCES
    ) or bool(
        (os.getenv("MIFACT_USUARIO", "").strip() or os.getenv("Mifact_usuario", "").strip())
        and (os.getenv("MIFACT_PASSWORD", "").strip() or os.getenv("Mifact_password", "").strip())
    )
    return jsonify(
        {
            "status": "ok",
            "xafiro_configurado": credentials_ready,
            "neo_configurado": neo_ready,
            "culqi_configurado": culqi_ready,
            "mifact_configurado": mifact_ready,
        }
    )


@app.post("/api/xafiro/export")
@permission_required("reportes.exportar")
def export_xafiro():
    payload = request.get_json(silent=True) or {}

    try:
        start_date = parse_iso_date(payload.get("fecha_inicio"), "La fecha inicial")
        end_date = parse_iso_date(payload.get("fecha_fin"), "La fecha final")
        if start_date > end_date:
            raise ValueError("La fecha inicial no puede ser posterior a la fecha final.")
        if (end_date - start_date).days > 366:
            raise ValueError("El rango máximo permitido es de 367 días.")

        content, filename, mimetype = download_xafiro_report(start_date, end_date)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except XafiroError as exc:
        return jsonify({"error": str(exc)}), 502

    return send_file(
        io.BytesIO(content),
        mimetype=mimetype,
        as_attachment=True,
        download_name=filename,
        max_age=0,
    )


@app.post("/api/neo/export")
@permission_required("reportes.exportar")
def export_neo():
    payload = request.get_json(silent=True) or {}

    try:
        start_date = parse_iso_date(payload.get("fecha_inicio"), "La fecha inicial")
        end_date = parse_iso_date(payload.get("fecha_fin"), "La fecha final")
        if start_date > end_date:
            raise ValueError("La fecha inicial no puede ser posterior a la fecha final.")
        if (end_date - start_date).days > 366:
            raise ValueError("El rango máximo permitido es de 367 días.")

        content, filename, mimetype = download_neo_report(start_date, end_date)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except NeoError as exc:
        return jsonify({"error": str(exc)}), 502

    return send_file(
        io.BytesIO(content),
        mimetype=mimetype,
        as_attachment=True,
        download_name=filename,
        max_age=0,
    )


@app.post("/api/culqi/export")
@permission_required("reportes.exportar")
def export_culqi():
    payload = request.get_json(silent=True) or {}

    try:
        start_date = parse_iso_date(payload.get("fecha_inicio"), "La fecha inicial")
        end_date = parse_iso_date(payload.get("fecha_fin"), "La fecha final")
        if start_date > end_date:
            raise ValueError("La fecha inicial no puede ser posterior a la fecha final.")
        if (end_date - start_date).days > 366:
            raise ValueError("El rango máximo permitido es de 367 días.")

        content, filename, mimetype = download_culqi_report(start_date, end_date)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except CulqiError as exc:
        return jsonify({"error": str(exc)}), 502

    return send_file(
        io.BytesIO(content),
        mimetype=mimetype,
        as_attachment=True,
        download_name=filename,
        max_age=0,
    )



@app.post("/api/consolidado/procesar")
@permission_required("reportes.exportar")
def procesar_consolidado():
    consolidation_type = normalize_text(request.form.get("tipo_consolidado", "resort"))
    try:
        content, summary = process_consolidado_uploads(request.files, consolidation_type)
    except ConsolidadoError as exc:
        return jsonify({"error": str(exc)}), 400

    response = send_file(
        io.BytesIO(content),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"consolidado_{consolidation_type}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        max_age=0,
    )
    response.headers["X-Consolidado-Summary"] = quote(json.dumps(summary, ensure_ascii=False))
    return response


@app.post("/api/mifact/export")
@permission_required("reportes.exportar")
def export_mifact():
    payload = request.get_json(silent=True) or {}

    try:
        start_date = parse_iso_date(payload.get("fecha_inicio"), "La fecha inicial")
        end_date = parse_iso_date(payload.get("fecha_fin"), "La fecha final")
        if start_date > end_date:
            raise ValueError("La fecha inicial no puede ser posterior a la fecha final.")
        if (end_date - start_date).days > 366:
            raise ValueError("El rango máximo permitido es de 367 días.")

        content, filename, mimetype = download_mifact_report(start_date, end_date)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except MifactError as exc:
        return jsonify({"error": str(exc)}), 502

    return send_file(
        io.BytesIO(content),
        mimetype=mimetype,
        as_attachment=True,
        download_name=filename,
        max_age=0,
    )


@app.errorhandler(413)
def payload_too_large(_error):
    return jsonify({"error": "La solicitud es demasiado grande."}), 413


@app.errorhandler(403)
def forbidden(_error):
    if request.path.startswith("/api/"):
        return jsonify({"error": "No autorizado."}), 403
    return render_template(
        "error.html",
        title="No autorizado",
        message="No tienes permisos para acceder a esta seccion.",
    ), 403


if __name__ == "__main__":
    host = os.getenv("APP_HOST", "127.0.0.1")
    port = int(os.getenv("APP_PORT", "5050"))
    debug = os.getenv("APP_DEBUG", "false").lower() == "true"
    app.run(host=host, port=port, debug=debug)
