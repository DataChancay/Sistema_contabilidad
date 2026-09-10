import io
import unittest
from datetime import date
from unittest.mock import patch

import requests
from openpyxl import Workbook, load_workbook
from werkzeug.datastructures import FileStorage

from app import app
from app import build_culqi_workbook
from app import build_mifact_workbook
from app import process_consolidado_uploads
from app import build_neo_workbook
from app import CulqiError
from app import culqi_request_error
from app import download_culqi_generated_workbook
from app import download_culqi_link_workbook
from app import mifact_request_error
from app import neo_request_error
from app import xafiro_request_error


def workbook_content(rows):
    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(row)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class FakeCulqiResponse:
    def __init__(self, payload=None, content=b"", status_code=200):
        self._payload = payload
        self.content = content or (b"{}" if payload is not None else b"")
        self.status_code = status_code
        self.text = ""

    def json(self):
        if self._payload is None:
            raise ValueError("No JSON payload")
        return self._payload

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError()


class AppRoutesTestCase(unittest.TestCase):
    def setUp(self):
        app.config.update(TESTING=True, AUTH_REQUIRED=False, WTF_CSRF_ENABLED=False)
        self.client = app.test_client()

    def test_index_is_available(self):
        response = self.client.get("/")

        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Contabilidad Chancay", response.data)
        self.assertIn(b"source-culqi", response.data)
        self.assertIn(b"source-mifact", response.data)
        self.assertIn(b"carga.json", response.data)
        self.assertIn(b"check.json", response.data)

    def test_animation_files_are_available(self):
        loading = self.client.get("/animaciones/carga.json")
        success = self.client.get("/animaciones/check.json")

        self.assertEqual(loading.status_code, 200)
        self.assertEqual(success.status_code, 200)
        self.assertIn(b"Download-lottie", loading.data)
        loading.close()
        success.close()

    def test_rejects_inverted_date_range(self):
        response = self.client.post(
            "/api/xafiro/export",
            json={"fecha_inicio": "2026-08-10", "fecha_fin": "2026-08-01"},
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("posterior", response.get_json()["error"])

    def test_rejects_invalid_date_format(self):
        response = self.client.post(
            "/api/xafiro/export",
            json={"fecha_inicio": "01/08/2026", "fecha_fin": "2026-08-10"},
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("AAAA-MM-DD", response.get_json()["error"])

    def test_explains_windows_socket_block(self):
        error = xafiro_request_error(
            requests.ConnectionError("[WinError 10013] Intento de acceso a un socket no permitido"),
            "iniciar sesión",
        )

        self.assertIn("bloqueó", str(error))
        self.assertIn("firewall", str(error))

    def test_explains_neo_windows_socket_block(self):
        error = neo_request_error(
            requests.ConnectionError("[WinError 10013] Intento de acceso a un socket no permitido"),
            "iniciar sesión",
        )

        self.assertIn("Neo", str(error))
        self.assertIn("firewall", str(error))

    def test_explains_culqi_windows_socket_block(self):
        error = culqi_request_error(
            requests.ConnectionError("[WinError 10013] Intento de acceso a un socket no permitido"),
            "iniciar sesión",
        )

        self.assertIn("Culqi", str(error))
        self.assertIn("firewall", str(error))

    def test_explains_mifact_windows_socket_block(self):
        error = mifact_request_error(
            requests.ConnectionError("[WinError 10013] Intento de acceso a un socket no permitido"),
            "descargar documentos",
        )

        self.assertIn("Mifact", str(error))
        self.assertIn("firewall", str(error))

    @patch("app.download_xafiro_report")
    def test_returns_report_as_attachment(self, download_report):
        download_report.return_value = (
            b"PK\x03\x04test-report",
            "facturacion_xafiro_2026-08-01_2026-08-02.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            "/api/xafiro/export",
            json={"fecha_inicio": "2026-08-01", "fecha_fin": "2026-08-02"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers["Content-Disposition"])
        self.assertTrue(response.data.startswith(b"PK\x03\x04"))

    @patch("app.download_neo_report")
    def test_returns_neo_report_as_attachment(self, download_report):
        download_report.return_value = (
            b"PK\x03\x04test-report",
            "facturacion_neo_2026-08-01_2026-08-02.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            "/api/neo/export",
            json={"fecha_inicio": "2026-08-01", "fecha_fin": "2026-08-02"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("facturacion_neo_2026-08-01_2026-08-02.xlsx", response.headers["Content-Disposition"])
        self.assertTrue(response.data.startswith(b"PK\x03\x04"))

    @patch("app.download_culqi_report")
    def test_returns_culqi_report_as_attachment(self, download_report):
        download_report.return_value = (
            b"PK\x03\x04test-report",
            "facturacion_culqi_2026-08-01_2026-08-02.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            "/api/culqi/export",
            json={"fecha_inicio": "2026-08-01", "fecha_fin": "2026-08-02"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("facturacion_culqi_2026-08-01_2026-08-02.xlsx", response.headers["Content-Disposition"])
        self.assertTrue(response.data.startswith(b"PK\x03\x04"))

    @patch("app.download_mifact_report")
    def test_returns_mifact_report_as_attachment(self, download_report):
        download_report.return_value = (
            b"PK\x03\x04test-report",
            "facturacion_mifact_2026-08-01_2026-08-02.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            "/api/mifact/export",
            json={"fecha_inicio": "2026-08-01", "fecha_fin": "2026-08-02"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("facturacion_mifact_2026-08-01_2026-08-02.xlsx", response.headers["Content-Disposition"])
        self.assertTrue(response.data.startswith(b"PK\x03\x04"))

    @patch("app.process_consolidado_uploads")
    def test_consolidado_route_uses_association_mode_and_filename(self, process_uploads):
        process_uploads.return_value = (b"PK\x03\x04association-report", {"conciliados_mifact": 1})

        response = self.client.post(
            "/api/consolidado/procesar",
            data={"tipo_consolidado": "asociacion"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(process_uploads.call_args.args[1], "asociacion")
        self.assertIn("consolidado_asociacion_", response.headers["Content-Disposition"])


    def test_processes_consolidado_with_xafiro_and_mifact_matches(self):
        culqi_headers = [f"Col {index}" for index in range(1, 27)]
        culqi_link_row = [None] * 26
        culqi_link_row[8] = "04/09/2026 10:32:15"
        culqi_link_row[10] = "Juan Carlos"
        culqi_link_row[11] = "Pérez Soto"
        culqi_link_row[25] = "S/ 120.50"
        culqi_full_row = [None] * 26
        culqi_full_row[8] = "04/09/2026"
        culqi_full_row[10] = "Ana"
        culqi_full_row[11] = "Torres"
        culqi_full_row[25] = 75

        xafiro_row = [None] * 18
        xafiro_row[2] = "B001"
        xafiro_row[3] = "42"
        xafiro_row[7] = "04/09/2026"
        xafiro_row[10] = "76543210"
        xafiro_row[11] = "PEREZ SOTO JUAN CARLOS"
        xafiro_row[17] = 120.5
        mifact_row = [None] * 17
        mifact_row[0] = "2026-09-04"
        mifact_row[3] = "F001"
        mifact_row[4] = "9"
        mifact_row[6] = "12345678"
        mifact_row[16] = "75.00"

        files = {
            "xafiro": (io.BytesIO(workbook_content([[None] * 18, xafiro_row])), "xafiro.xlsx"),
            "mifact": (io.BytesIO(workbook_content([[None] * 17, mifact_row])), "mifact.xlsx"),
            "culqilink": (io.BytesIO(workbook_content([culqi_headers, culqi_link_row])), "culqilink.xlsx"),
            "culqifull": (io.BytesIO(workbook_content([culqi_headers, culqi_full_row])), "culqifull.xlsx"),
        }

        content, summary = process_consolidado_uploads({key: FileStorage(stream=value, filename=name) for key, (value, name) in files.items()})
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))

        self.assertEqual(summary["total_operaciones_culqi"], 2)
        self.assertEqual(summary["conciliados_xafiro"], 1)
        self.assertEqual(summary["conciliados_mifact"], 1)
        self.assertIn("Origen comprobante", rows[0])
        self.assertEqual(rows[1][-2:], ("Xafiro", "CONCILIADO"))
        self.assertEqual(rows[2][-2:], ("Mifact", "CONCILIADO"))

    def test_processes_consolidado_places_rejected_culqi_rows_first_in_red(self):
        culqi_headers = [f"Col {index}" for index in range(1, 38)]
        rejected_row = [None] * 37
        rejected_row[8] = "04/09/2026"
        rejected_row[25] = 120.5
        rejected_row[36] = "rechazada"
        valid_row = [None] * 37
        valid_row[8] = "05/09/2026"
        valid_row[25] = 75
        valid_row[36] = "exitosa"
        annulled_row = [None] * 37
        annulled_row[8] = "06/09/2026"
        annulled_row[25] = 90
        annulled_row[36] = "anulada"

        xafiro_row = [None] * 18
        xafiro_row[2] = "B001"
        xafiro_row[3] = "42"
        xafiro_row[7] = "04/09/2026"
        xafiro_row[10] = "76543210"
        xafiro_row[11] = "CLIENTE RECHAZADO"
        xafiro_row[17] = 120.5
        mifact_row = [None] * 17
        mifact_row[0] = "05/09/2026"
        mifact_row[3] = "F001"
        mifact_row[4] = "9"
        mifact_row[6] = "12345678"
        mifact_row[16] = "75.00"

        files = {
            "xafiro": (io.BytesIO(workbook_content([[None] * 18, xafiro_row])), "xafiro.xlsx"),
            "mifact": (io.BytesIO(workbook_content([[None] * 17, mifact_row])), "mifact.xlsx"),
            "culqilink": (io.BytesIO(workbook_content([culqi_headers, rejected_row])), "culqilink.xlsx"),
            "culqifull": (io.BytesIO(workbook_content([culqi_headers, valid_row, annulled_row])), "culqifull.xlsx"),
        }

        content, summary = process_consolidado_uploads({key: FileStorage(stream=value, filename=name) for key, (value, name) in files.items()})
        workbook = load_workbook(io.BytesIO(content))
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(min_row=2))

        self.assertEqual(summary["total_operaciones_culqi"], 3)
        self.assertEqual(summary["conciliados_xafiro"], 0)
        self.assertEqual(summary["conciliados_mifact"], 1)
        self.assertEqual([row[-1].value for row in rows], ["RECHAZADA", "ANULADA", "CONCILIADO"])
        for row in rows[:2]:
            self.assertTrue(all((cell.font.color and cell.font.color.rgb or "").endswith("FF0000") for cell in row))
        self.assertEqual(rows[2][-2].value, "Mifact")

    def test_processes_asociacion_without_xafiro_and_keeps_same_matching_rules(self):
        culqi_headers = [f"Col {index}" for index in range(1, 38)]
        mifact_match = [None] * 37
        mifact_match[8] = "05/09/2026"
        mifact_match[25] = 75
        xafiro_only_match = [None] * 37
        xafiro_only_match[8] = "06/09/2026"
        xafiro_only_match[25] = 120.5
        rejected_row = [None] * 37
        rejected_row[8] = "07/09/2026"
        rejected_row[25] = 40
        rejected_row[36] = "rechazada"

        mifact_row = [None] * 17
        mifact_row[0] = "05/09/2026"
        mifact_row[3] = "F001"
        mifact_row[4] = "9"
        mifact_row[6] = "12345678"
        mifact_row[16] = 75

        files = {
            "mifact": FileStorage(stream=io.BytesIO(workbook_content([[None] * 17, mifact_row])), filename="mifact.xlsx"),
            "culqilink": FileStorage(stream=io.BytesIO(workbook_content([culqi_headers, mifact_match])), filename="culqilink.xlsx"),
            "culqifull": FileStorage(stream=io.BytesIO(workbook_content([culqi_headers, xafiro_only_match, rejected_row])), filename="culqifull.xlsx"),
        }

        content, summary = process_consolidado_uploads(files, "asociacion")
        workbook = load_workbook(io.BytesIO(content))
        rows = list(workbook.active.iter_rows(min_row=2))

        self.assertEqual(summary["conciliados_xafiro"], 0)
        self.assertEqual(summary["conciliados_mifact"], 1)
        self.assertEqual([row[-1].value for row in rows], ["RECHAZADA", "NO ENCONTRADO", "CONCILIADO"])
        self.assertTrue(all((cell.font.color and cell.font.color.rgb or "").endswith("FF0000") for cell in rows[0]))
        self.assertEqual(rows[-1][-2].value, "Mifact")

    def test_processes_consolidado_matches_balanced_duplicate_amounts(self):
        culqi_headers = [f"Col {index}" for index in range(1, 27)]
        row_one = [None] * 26
        row_one[8] = "08/07/2026"
        row_one[25] = 300
        row_two = [None] * 26
        row_two[8] = "08/07/2026"
        row_two[25] = 300
        xafiro_first = [None] * 18
        xafiro_first[2] = "BX01"
        xafiro_first[3] = "11448"
        xafiro_first[7] = "08/07/2026"
        xafiro_first[10] = "42612952"
        xafiro_first[11] = "CHUCHON RAMIREZ WILDER"
        xafiro_first[17] = 300
        xafiro_second = [None] * 18
        xafiro_second[2] = "BX01"
        xafiro_second[3] = "11447"
        xafiro_second[7] = "08/07/2026"
        xafiro_second[10] = "45521743"
        xafiro_second[11] = "MOSCOSO BALLON KARINA STEFANY"
        xafiro_second[17] = 300
        mifact_row = [None] * 17
        mifact_row[0] = "09/07/2026"
        mifact_row[16] = 10

        files = {
            "xafiro": (io.BytesIO(workbook_content([[None] * 18, xafiro_first, xafiro_second])), "xafiro.xlsx"),
            "mifact": (io.BytesIO(workbook_content([[None] * 17, mifact_row])), "mifact.xlsx"),
            "culqilink": (io.BytesIO(workbook_content([culqi_headers, row_one])), "culqilink.xlsx"),
            "culqifull": (io.BytesIO(workbook_content([culqi_headers, row_two])), "culqifull.xlsx"),
        }

        content, summary = process_consolidado_uploads({key: FileStorage(stream=value, filename=name) for key, (value, name) in files.items()})
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))

        self.assertEqual(summary["conciliados_xafiro"], 2)
        self.assertEqual(summary["registros_revisar"], 0)
        self.assertEqual([row[-1] for row in rows], ["CONCILIADO", "CONCILIADO"])
        self.assertEqual([row[-4] for row in rows], ["11448", "11447"])


    def test_processes_consolidado_matches_gateway_count_when_more_invoices_exist(self):
        culqi_headers = [f"Col {index}" for index in range(1, 27)]
        culqi_rows = []
        for _index in range(4):
            row = [None] * 26
            row[8] = "08/07/2026"
            row[25] = 20
            culqi_rows.append(row)

        xafiro_rows = []
        for index in range(5):
            row = [None] * 18
            row[2] = "BX01"
            row[3] = str(200 + index)
            row[7] = "08/07/2026"
            row[10] = str(45000000 + index)
            row[11] = f"CLIENTE {index}"
            row[17] = 20
            xafiro_rows.append(row)

        mifact_row = [None] * 17
        mifact_row[0] = "09/07/2026"
        mifact_row[16] = 10

        files = {
            "xafiro": (io.BytesIO(workbook_content([[None] * 18, *xafiro_rows])), "xafiro.xlsx"),
            "mifact": (io.BytesIO(workbook_content([[None] * 17, mifact_row])), "mifact.xlsx"),
            "culqilink": (io.BytesIO(workbook_content([culqi_headers, *culqi_rows[:2]])), "culqilink.xlsx"),
            "culqifull": (io.BytesIO(workbook_content([culqi_headers, *culqi_rows[2:]])), "culqifull.xlsx"),
        }

        content, summary = process_consolidado_uploads({key: FileStorage(stream=value, filename=name) for key, (value, name) in files.items()})
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(min_row=2, values_only=True))

        self.assertEqual(summary["conciliados_xafiro"], 4)
        self.assertEqual(summary["registros_revisar"], 0)
        self.assertEqual([row[-1] for row in rows], ["CONCILIADO"] * 4)
        self.assertEqual([row[-4] for row in rows], ["200", "201", "202", "203"])

    def test_processes_consolidado_marks_mifact_ambiguity_for_review(self):
        culqi_headers = [f"Col {index}" for index in range(1, 27)]
        row_one = [None] * 26
        row_one[8] = "04/09/2026"
        row_one[25] = 90
        row_two = [None] * 26
        row_two[8] = "04/09/2026"
        row_two[25] = 90
        xafiro_row = [None] * 18
        xafiro_row[7] = "05/09/2026"
        xafiro_row[17] = 10
        mifact_row = [None] * 17
        mifact_row[0] = "04/09/2026"
        mifact_row[16] = 90

        files = {
            "xafiro": (io.BytesIO(workbook_content([[None] * 18, xafiro_row])), "xafiro.xlsx"),
            "mifact": (io.BytesIO(workbook_content([[None] * 17, mifact_row])), "mifact.xlsx"),
            "culqilink": (io.BytesIO(workbook_content([culqi_headers, row_one])), "culqilink.xlsx"),
            "culqifull": (io.BytesIO(workbook_content([culqi_headers, row_two])), "culqifull.xlsx"),
        }

        content, summary = process_consolidado_uploads({key: FileStorage(stream=value, filename=name) for key, (value, name) in files.items()})
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        states = [row[-1] for row in worksheet.iter_rows(min_row=2, values_only=True)]

        self.assertEqual(summary["registros_revisar"], 2)
        self.assertEqual(states, ["REVISAR", "REVISAR"])

    def test_builds_neo_workbook(self):
        content = build_neo_workbook(
            [
                {
                    "fecha_operacion": "2026-08-01",
                    "producto": "Ticket adulto",
                    "cantidad": 2,
                    "precio": 30,
                }
            ]
        )

        self.assertTrue(content.startswith(b"PK\x03\x04"))

    def test_builds_culqi_workbook(self):
        headers = ["Empresa", "Comercio", "Producto", "ID Venta"]
        full = workbook_content(
            [
                headers,
                ["Empresa 1", "Comercio full", "CulqiFull", "full-1"],
            ]
        )
        link = workbook_content(
            [
                headers,
                ["Empresa 1", "Comercio link", "CulqiLink", "link-1"],
            ]
        )

        content = build_culqi_workbook([("CulqiFull", full), ("CulqiLink", link)])
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active

        self.assertEqual([cell.value for cell in worksheet[1]], headers)
        self.assertEqual(
            [cell.value for cell in worksheet[2]],
            ["Empresa 1", "Comercio full", "CulqiFull", "full-1"],
        )
        self.assertEqual(
            [cell.value for cell in worksheet[3]],
            ["Empresa 1", "Comercio link", "CulqiLink", "link-1"],
        )
        self.assertTrue(content.startswith(b"PK\x03\x04"))

    def test_rejects_culqi_workbooks_with_different_headers(self):
        full = workbook_content([["Empresa", "Comercio"], ["Empresa 1", "Full"]])
        link = workbook_content([["Empresa", "Producto"], ["Empresa 1", "Link"]])

        with self.assertRaises(CulqiError):
            build_culqi_workbook([("CulqiFull", full), ("CulqiLink", link)])

    def test_checks_culqi_generated_report_with_post_and_merchant_ids(self):
        class FakeSession:
            def __init__(self):
                self.requests = []

            def request(self, method, url, timeout=None, **kwargs):
                self.requests.append((method, url, kwargs))
                if url.endswith("/sales/download-report"):
                    return FakeCulqiResponse({"success": True, "data": {"id": "report-1"}})
                if url.endswith("/sales/get-report/report-1"):
                    return FakeCulqiResponse(
                        {"success": True, "data": "https://files.culqi.test/report.xlsx"}
                    )
                return FakeCulqiResponse({"success": False}, status_code=404)

            def get(self, url, timeout=None):
                self.requests.append(("GET_FILE", url, {}))
                return FakeCulqiResponse(content=b"PK\x03\x04xlsx")

        session = FakeSession()

        content = download_culqi_generated_workbook(
            session,
            "sales/download-report",
            "sales/get-report",
            {"merchantIds": ["200000000180243"]},
            "CulqiFull",
            merchant_ids=["200000000180243"],
        )

        self.assertEqual(content, b"PK\x03\x04xlsx")
        self.assertEqual(session.requests[1][0], "POST")
        self.assertTrue(session.requests[1][1].endswith("/sales/get-report/report-1"))
        self.assertEqual(
            session.requests[1][2]["json"],
            {"merchantIds": ["200000000180243"]},
        )

    @patch("app.download_culqi_generated_workbook", return_value=b"PK\x03\x04report")
    def test_downloads_culqilink_from_sales_report(self, download_generated):
        class FakeSession:
            def __init__(self):
                self.headers = {}

        session = FakeSession()
        merchant = {
            "id": "parent-1",
            "merchantId": "200000000180243",
            "products": [{"id": "7", "name": "CulqiLink", "publicKey": "pk_live_test"}],
        }

        content = download_culqi_link_workbook(
            session,
            merchant,
            date(2026, 7, 1),
            date(2026, 7, 31),
        )

        self.assertEqual(content, b"PK\x03\x04report")
        self.assertEqual(download_generated.call_args.args[1], "sales/download-report")
        self.assertEqual(download_generated.call_args.args[2], "sales/get-report")
        self.assertEqual(download_generated.call_args.args[3]["productTypes"], [7])
        self.assertEqual(download_generated.call_args.args[3]["merchantStatusList"], download_generated.call_args.args[3]["statusList"])
        self.assertEqual(download_generated.call_args.args[3]["currencyList"], ["PEN", "USD"])
        self.assertEqual(download_generated.call_args.args[3]["cardTypeList"], ["CARD", "EF"])
        self.assertEqual(download_generated.call_args.args[3]["cardBrandList"], ["04", "05", "03", "07", "08", "09", "99", "97"])
        self.assertEqual(download_generated.call_args.args[3]["partialPaymentList"], ["CR", "DB", "PRE"])
        self.assertEqual(download_generated.call_args.args[3]["walletList"], ["01", "02", "03"])
        self.assertEqual(download_generated.call_args.kwargs["merchant_ids"], ["200000000180243"])
        self.assertEqual(
            download_generated.call_args.args[3]["statusList"],
            ["01", "02", "03", "04", "05", "06", "07", "08", "09", "10", "11", "12"],
        )
        self.assertNotIn("cardNumber", download_generated.call_args.args[3])
        self.assertNotIn("search", download_generated.call_args.args[3])
        self.assertNotIn("pkLives", download_generated.call_args.args[3])

    def test_builds_culqi_workbook_with_empty_real_headers(self):
        content = build_culqi_workbook(
            [
                (
                    "CulqiFull",
                    workbook_content([["Empresa", "Comercio", "Producto", "ID Venta"]]),
                ),
            ]
        )

        self.assertTrue(content.startswith(b"PK\x03\x04"))

    def test_builds_mifact_workbook_without_renaming_headers(self):
        first = workbook_content(
            [
                ["Fecha Emisión", "Tipo Documento", "Total"],
                ["2026-08-01", "Boleta de venta", 20],
            ]
        )
        second = workbook_content(
            [
                ["Fecha Emisión", "Tipo Documento", "Total"],
                ["2026-08-02", "Factura", 35],
            ]
        )

        content = build_mifact_workbook([("Empresa 1", first), ("Empresa 2", second)])

        self.assertTrue(content.startswith(b"PK\x03\x04"))


if __name__ == "__main__":
    unittest.main()
